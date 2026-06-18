#!/usr/bin/env python3
"""엑셀(01_매체마스터) → media_image 테이블 적재 (database-design.md §2.4).

detail.mediaItemImages(JSON 배열)를 URL 단위로 분해한다.
  image_url   = entry.fullPath
  sort_order  = 배열 내 순서
  is_thumbnail = entry.imageType == 'THUMBNAIL'
media 행이 먼저 존재해야 한다(FK). 적재 전 기존 media_image 를 비운다.

사용법:
    python -m scripts.import_media_images <엑셀파일>
"""
from __future__ import annotations

import json
import sys
import uuid
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

load_dotenv()

from src.config import get_settings
from src.models.media_image import MediaImage
from src.models.media_master import Media

SHEET = "01_매체마스터"

engine = create_engine(get_settings().database_url, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine)


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() == "none":
        return None
    return s


def _resolve_idx(ws) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(h).strip(): i for i, h in enumerate(header_row) if h is not None}


def _images_from_row(media_id: str, raw: str | None) -> list[dict]:
    s = _clean(raw)
    if s is None:
        return []
    try:
        arr = json.loads(s)
    except (json.JSONDecodeError, ValueError):
        return []
    out: list[dict] = []
    order = 0
    for entry in arr:
        if not isinstance(entry, dict):
            continue
        url = _clean(entry.get("fullPath"))
        if url is None:
            continue
        out.append(
            dict(
                id=uuid.uuid4(),
                media_id=media_id,
                image_url=url,
                sort_order=order,
                is_thumbnail=(entry.get("imageType") == "THUMBNAIL"),
            )
        )
        order += 1
    return out


def import_images(wb, session: Session, replace: bool = True) -> tuple[int, int]:
    if replace:
        session.query(MediaImage).delete()
        session.flush()
    valid_ids = {m[0] for m in session.query(Media.media_id).all()}
    ws = wb[SHEET]
    idx = _resolve_idx(ws)
    mi_col = idx.get("detail.mediaItemImages")
    id_col = idx.get("media_id")
    rows_count = 0
    media_with_images = 0
    missing_media = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        media_id = _clean(row[id_col]) if id_col is not None else None
        if media_id is None:
            continue
        if media_id not in valid_ids:
            missing_media += 1
            continue
        raw = row[mi_col] if mi_col is not None and mi_col < len(row) else None
        imgs = _images_from_row(media_id, raw)
        if not imgs:
            continue
        media_with_images += 1
        for data in imgs:
            session.add(MediaImage(**data))
            rows_count += 1
        if rows_count % 500 == 0:
            session.flush()
    session.commit()
    if missing_media:
        print(f"  [skip] media 행 없는 media_id {missing_media}건")
    return rows_count, media_with_images


def main(path: str) -> None:
    p = Path(path)
    print(f"[import] 파일: {p}")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"[error] 시트 '{SHEET}' 없음. 보유: {wb.sheetnames}")
        sys.exit(1)
    session = SessionLocal()
    try:
        rows, media_n = import_images(wb, session)
        print(f"[import] media_image 적재: {rows}건 (매체 {media_n}개)")
        print(f"[검증] media_image count: {session.query(MediaImage).count()}건")
    finally:
        session.close()
    print("[완료]")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python -m scripts.import_media_images <엑셀파일>")
        sys.exit(1)
    main(sys.argv[1])
