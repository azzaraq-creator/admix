#!/usr/bin/env python3
"""엑셀(01_매체마스터) → media 테이블 적재 (database-design.md §2.1).

시트 01_매체마스터: 행1=영문헤더, 행2=한글헤더, 행3+=데이터(913 매체).
영문 헤더로 컬럼 인덱스를 찾아 media 단일값 컬럼만 적재한다.
plan/prop/유동인구/라벨/이미지 등 반복·자식 데이터는 적재 대상이 아니다.

사용법:
    python -m scripts.import_media_master <엑셀파일>
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

load_dotenv()

from src.config import get_settings
from src.models.media_master import Media

SHEET = "01_매체마스터"

engine = create_engine(get_settings().database_url, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine)


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() == "none":
        return None
    return s


def _int(v) -> int | None:
    s = _clean(v)
    if s is None:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return int(float(s))
        except ValueError:
            return None


def _num(v) -> Decimal | None:
    s = _clean(v)
    if s is None:
        return None
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _bool_yn(v) -> bool | None:
    s = _clean(v)
    if s is None:
        return None
    s = s.upper()
    if s in ("Y", "TRUE", "1"):
        return True
    if s in ("N", "FALSE", "0"):
        return False
    return None


def _json(v):
    s = _clean(v)
    if s is None:
        return None
    try:
        return json.loads(s)
    except (json.JSONDecodeError, ValueError):
        return None


def _labels(v) -> list[str] | None:
    s = _clean(v)
    if s is None:
        return None
    parts = [p.strip() for p in s.split("|") if p.strip()]
    return parts or None


def _ts(v) -> datetime | None:
    s = _clean(v)
    if s is None:
        return None
    s = s.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


def _resolve_idx(ws) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(h).strip(): i for i, h in enumerate(header_row) if h is not None}


def _row_to_media(row, idx: dict[str, int]) -> dict | None:
    def g(name: str):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    media_id = _clean(g("media_id"))
    if media_id is None:
        return None

    map_bounds = {
        "northEastLatitude": _num(g("detail.mapBoundsVo.northEastLatitude")),
        "northEastLongitude": _num(g("detail.mapBoundsVo.northEastLongitude")),
        "southWestLatitude": _num(g("detail.mapBoundsVo.southWestLatitude")),
        "southWestLongitude": _num(g("detail.mapBoundsVo.southWestLongitude")),
    }
    map_bounds = {k: float(v) for k, v in map_bounds.items() if v is not None} or None

    return dict(
        media_id=media_id,
        source_detail_id=_int(g("detail.id")),
        name=_clean(g("detail.name")),
        second_name=_clean(g("detail.secondName")),
        building_name=_clean(g("detail.buildingName")),
        category_small=_clean(g("detail.mediaItemCategory.displayValue")),
        category_large=_clean(g("detail.mediaItemCategory.parentCategory.displayValue")),
        category_id=_int(g("detail.mediaItemCategory.id")),
        parent_category_code=_clean(g("detail.mediaItemCategory.parentCategory.code")),
        ooh_type=_clean(g("detail.oohType")),
        exposure_type=_clean(g("detail.exposureType")),
        sales_type=_clean(g("sales_type")),
        media_source=_clean(g("media_source")),
        loc_code=_clean(g("market_loc_code")),
        loc_label=_clean(g("LOC (지역 라벨)")),
        market_area=_clean(g("market_area")),
        market_dong=_clean(g("market_dong")),
        legal_dong=_clean(g("legal_dong")),
        accurate_address=_clean(g("accurate_address")),
        full_address_jibun=_clean(g("full_address_jibun")),
        address=_clean(g("detail.address")),
        address_detail=_clean(g("detail.addressDetail")),
        district=_clean(g("detail.district")),
        city=_clean(g("detail.city")),
        latitude=_num(g("detail.latitude")),
        longitude=_num(g("detail.longitude")),
        market_keyword=_clean(g("market_keyword")),
        moving_location_detail=_clean(g("detail.movingLocationDetail")),
        audience_summary=_clean(g("audience_summary")),
        min_advertisement_fee_krw=_int(g("min_advertisement_fee_krw")),
        max_advertisement_fee_krw=_int(g("max_advertisement_fee_krw")),
        min_production_fee_krw=_int(g("min_production_fee_krw")),
        max_production_fee_krw=_int(g("max_production_fee_krw")),
        production_fees_summary=_clean(g("production_fees_summary")),
        any_production_fee_yn=_bool_yn(g("any_production_fee_yn")),
        plan_count=_int(g("plan_count")),
        execution_status=_clean(g("execution_status")),
        lead_time_bizdays=_int(g("lead_time_bizdays")),
        device_quantity=_int(g("detail.deviceQuantity")),
        surface_quantity=_int(g("detail.surfaceQuantity")),
        media_shape=_clean(g("detail.mediaShape")),
        media_shape_summary=_clean(g("detail.mediaShapeSummary")),
        properties_count=_int(g("properties_count")),
        properties_summary=_clean(g("properties_summary")),
        properties_extra_json=_json(g("properties_extra_json")),
        final_grade=_clean(g("final_grade")),
        feature=_clean(g("feature")),
        quality_score=_num(g("quality_score")),
        gangnam_dong_grade=_clean(g("gangnam_dong_grade")),
        gangnam_grade_reason=_clean(g("gangnam_grade_reason")),
        grade_method=_clean(g("grade_method")),
        grade_evidence=_clean(g("grade_evidence")),
        area_evidence=_clean(g("area_evidence")),
        ind_evidence=_clean(g("ind_evidence")),
        thumbnail_url=_clean(g("thumbnail_url")),
        image_count=_int(g("image_count")),
        is_newly_built_yn=_bool_yn(g("detail.isNewlyBuiltYn")),
        is_popular_yn=_bool_yn(g("detail.isPopularYn")),
        popular_type=_clean(g("detail.popularType")),
        special_remarks=_clean(g("detail.specialRemarks")),
        children_count=_int(g("detail.childrenCount")),
        device_type=_clean(g("detail.deviceType")),
        description=_clean(g("detail.description")),
        markers_vo=_json(g("detail.markersVo")),
        properties_type=_clean(g("detail.propertiesType")),
        road_view_heading=_num(g("detail.roadViewHeading")),
        road_view_latitude=_num(g("detail.roadViewLatitude")),
        road_view_longitude=_num(g("detail.roadViewLongitude")),
        road_view_pitch=_num(g("detail.roadViewPitch")),
        map_bounds_vo=map_bounds,
        recommended_media_items=_json(g("detail.recommendedMediaItems")),
        list_labels=_labels(g("list.labels")),
        company_media_id=_int(g("detail.companyMediaId")),
        company_mapper_user_id=_int(g("detail.companyMapperUserId")),
        source_created_at=_ts(g("detail.createdAt")),
        source_updated_at=_ts(g("detail.updatedAt")),
    )


def import_media(wb, session: Session, replace: bool = True) -> int:
    if replace:
        session.query(Media).delete()
        session.flush()
    ws = wb[SHEET]
    idx = _resolve_idx(ws)
    seen: set[str] = set()
    count = 0
    skipped_dup = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        data = _row_to_media(row, idx)
        if data is None:
            continue
        if data["media_id"] in seen:
            skipped_dup += 1
            continue
        seen.add(data["media_id"])
        session.add(Media(**data))
        count += 1
        if count % 200 == 0:
            session.flush()
    session.commit()
    if skipped_dup:
        print(f"  [skip] 중복 media_id {skipped_dup}건")
    return count


def main(path: str) -> None:
    p = Path(path)
    print(f"[import] 파일: {p}")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"[error] 시트 '{SHEET}' 없음. 보유: {wb.sheetnames}")
        sys.exit(1)
    session = SessionLocal()
    try:
        n = import_media(wb, session)
        print(f"[import] media 적재: {n}건")
        print(f"[검증] media count: {session.query(Media).count()}건")
    finally:
        session.close()
    print("[완료]")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python -m scripts.import_media_master <엑셀파일>")
        sys.exit(1)
    main(sys.argv[1])
