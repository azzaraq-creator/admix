"""엑셀 prop1~prop5 컬럼을 media.properties_extra_json 으로 조립·보완 적재.

원본 매체마스터 엑셀은 규격(width/height/unit)·수량을 prop1~prop5 반복그룹
컬럼으로 보관한다. 최초 import 는 (대부분 비어있는) properties_extra_json
단일 컬럼만 읽어 규격이 누락됐다. 이 스크립트가 prop1~5 를 JSON 배열로
조립해 properties_extra_json 을 채운다.

사용: python scripts/backfill_media_properties.py [엑셀경로]
"""
from __future__ import annotations

import json
import sys

import openpyxl
import psycopg2

DEFAULT_XLSX = "/Users/lala/Desktop/아우라웍스/ADMIX_최종본_v3.xlsx"
DB_DSN = "postgresql://postgres:postgres@localhost:5433/ooh_recommend"
SHEET = "01_매체마스터"
MAX_PROP = 5

# 출력 키 = 엑셀 prop{n}.{컬럼} 접미사
PROP_FIELDS = [
    "propertyKey",
    "displayValue",
    "propertyValue",
    "propertyUnit",
    "propertyWidthValue",
    "propertyHeightValue",
    "deviceQuantity",
    "surfaceQuantity",
]


def _norm(value):
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def main() -> None:
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb[SHEET]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    mid_i = idx["media_id"]

    updates: list[tuple[str, str]] = []
    # 1행=영문 헤더, 2행=한글 설명, 3행부터 실제 데이터
    for row in ws.iter_rows(min_row=3, values_only=True):
        media_id = _norm(row[mid_i])
        if not media_id:
            continue
        props = []
        for n in range(1, MAX_PROP + 1):
            entry = {}
            has_value = False
            for field in PROP_FIELDS:
                ci = idx.get(f"prop{n}.{field}")
                val = _norm(row[ci]) if ci is not None else None
                entry[field] = val
                if val is not None:
                    has_value = True
            if has_value:
                props.append(entry)
        if props:
            updates.append((media_id, json.dumps(props, ensure_ascii=False)))

    conn = psycopg2.connect(DB_DSN)
    cur = conn.cursor()
    rows = 0
    for media_id, payload in updates:
        cur.execute(
            "UPDATE media SET properties_extra_json = %s::jsonb WHERE media_id = %s",
            (payload, media_id),
        )
        rows += cur.rowcount
    conn.commit()
    cur.close()
    conn.close()
    print(f"prepared={len(updates)} rows_updated={rows}")


if __name__ == "__main__":
    main()
