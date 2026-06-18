#!/usr/bin/env python3
"""엑셀(01_매체마스터) → media_plan 테이블 적재 (database-design.md §2.2).

plan1~plan5 를 행으로 정규화. plan{n}.{field} 헤더명으로 매핑하므로
plan 별 컬럼 순서/누락(plan3~5 일부 필드 없음)과 무관하다.
plan 존재 판정: plan{n}.productName 비어있지 않음.
media 행이 먼저 존재해야 한다(FK). 적재 전 기존 media_plan 을 비운다.

사용법:
    python -m scripts.import_media_plans <엑셀파일>
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

load_dotenv()

from src.config import get_settings
from src.models.media_master import Media
from src.models.media_plan import MediaPlan

SHEET = "01_매체마스터"

engine = create_engine(get_settings().database_url, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine)


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() == "none":
        return None
    return s


def _int(v) -> int | None:
    s = _clean(v)
    if s is None:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return int(float(s))
        except ValueError:
            return None


def _bool_yn(v) -> bool | None:
    s = _clean(v)
    if s is None:
        return None
    s = s.upper()
    if s in ("Y", "TRUE", "1"):
        return True
    if s in ("N", "FALSE", "0"):
        return False
    return None


def _resolve_idx(ws) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(h).strip(): i for i, h in enumerate(header_row) if h is not None}


def _plan_from_row(row, idx: dict[str, int], media_id: str, n: int) -> dict | None:
    def g(field: str):
        k = f"plan{n}.{field}"
        i = idx.get(k)
        return row[i] if i is not None and i < len(row) else None

    if _clean(g("productName")) is None:
        return None

    return dict(
        media_id=media_id,
        plan_no=n,
        product_name=_clean(g("productName")),
        product_display_name=_clean(g("productDisplayName")),
        product_master_type=_clean(g("productMasterType")),
        contractual_duration=_int(g("contractualDuration")),
        contractual_duration_type=_clean(g("contractualDurationType")),
        advertisement_fee=_int(g("advertisementFee")),
        is_advertisement_fee_yn=_bool_yn(g("isAdvertisementFeeYn")),
        production_fee=_int(g("productionFee")),
        is_production_fee_yn=_bool_yn(g("isProductionFeeYn")),
        exposure_duration_seconds=_int(g("exposureDurationSeconds")),
        exposure_count=_int(g("exposureCount")),
        broadcasts_count_auto=_int(g("broadcastsCountAuto")),
        broadcasts_count_manual=_int(g("broadcastsCountManual")),
        ooh_type=_clean(g("oohType")),
        ooh_kind_type=_clean(g("oohKindType")),
        default_device_type=_clean(g("defaultDeviceType")),
        default_device_quantity=_int(g("defaultDeviceQuantity")),
        default_surface_quantity=_int(g("defaultSurfaceQuantity")),
        active_device_type=_clean(g("activeDeviceType")),
        active_device_quantity=_int(g("activeDeviceQuantity")),
        active_surface_quantity=_int(g("activeSurfaceQuantity")),
        pm_count=_int(g("pmCount")),
        operation_day_of_week=_clean(g("operationDayOfWeek")),
        operation_start_time=_clean(g("operationStartTime")),
        operation_end_time=_clean(g("operationEndTime")),
        operation_hours=_int(g("operationHours")),
    )


def import_plans(wb, session: Session, replace: bool = True) -> tuple[int, dict[int, int]]:
    if replace:
        session.query(MediaPlan).delete()
        session.flush()
    valid_ids = {m[0] for m in session.query(Media.media_id).all()}
    ws = wb[SHEET]
    idx = _resolve_idx(ws)
    id_col = idx.get("media_id")
    total = 0
    per_plan: dict[int, int] = {n: 0 for n in range(1, 6)}
    missing_media = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        media_id = _clean(row[id_col]) if id_col is not None else None
        if media_id is None:
            continue
        if media_id not in valid_ids:
            missing_media += 1
            continue
        for n in range(1, 6):
            data = _plan_from_row(row, idx, media_id, n)
            if data is None:
                continue
            session.add(MediaPlan(**data))
            per_plan[n] += 1
            total += 1
            if total % 500 == 0:
                session.flush()
    session.commit()
    if missing_media:
        print(f"  [skip] media 행 없는 media_id {missing_media}건")
    return total, per_plan


def main(path: str) -> None:
    p = Path(path)
    print(f"[import] 파일: {p}")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"[error] 시트 '{SHEET}' 없음. 보유: {wb.sheetnames}")
        sys.exit(1)
    session = SessionLocal()
    try:
        total, per_plan = import_plans(wb, session)
        print(f"[import] media_plan 적재: {total}건")
        for n in range(1, 6):
            print(f"    plan{n}: {per_plan[n]}건")
        print(f"[검증] media_plan count: {session.query(MediaPlan).count()}건")
    finally:
        session.close()
    print("[완료]")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python -m scripts.import_media_plans <엑셀파일>")
        sys.exit(1)
    main(sys.argv[1])
