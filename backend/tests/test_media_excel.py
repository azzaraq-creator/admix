"""어드민 엑셀 다운로드 — 폼과 같은 구조, 사용하지 않는 값 제외, 빠진 컬럼 없음."""
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import inspect as sa_inspect

from src.database import SessionLocal
from src.models.media_master import Media
from src.models.media_plan import MediaPlan
from src.services import kakao_address, media_excel

# 카카오 주소 검색 대신 쓰는 가짜 결과 — 테스트는 바깥 API 를 부르지 않는다.
_FAKE_ADDRESSES = {
    "서울 강남구 강남대로 396": {
        "accurate_address": "서울 강남구 강남대로 396 (역삼동)",
        "full_address_jibun": "서울 강남구 역삼동 858",
        "building_name": "강남역 빌딩",
        "legal_dong": "역삼동",
        "city": "서울",
        "district": "강남구",
        "latitude": 37.4979,
        "longitude": 127.0276,
    },
}


@pytest.fixture(autouse=True)
def fake_kakao_address(monkeypatch):
    calls: list[str] = []

    def lookup(address, client=None):
        calls.append(address)
        if address not in _FAKE_ADDRESSES:
            raise kakao_address.AddressLookupError(f"주소를 찾지 못했습니다: {address}")
        return dict(_FAKE_ADDRESSES[address])

    monkeypatch.setattr(kakao_address, "lookup", lookup)
    return calls


def test_every_column_is_exported_or_marked_unused():
    """새 컬럼을 추가하면 엑셀(SECTIONS)이나 UNUSED_COLUMNS 중 한쪽에 넣어야 한다."""
    exported = {key for _, cols in media_excel.SECTIONS for key, _, _ in cols}
    all_cols = {c.key for c in sa_inspect(Media).mapper.column_attrs}
    assert exported.isdisjoint(media_excel.UNUSED_COLUMNS)
    assert all_cols - exported - media_excel.UNUSED_COLUMNS == set()


def test_export_layout_and_values():
    db = SessionLocal()
    try:
        ws = load_workbook(BytesIO(media_excel.export_media_xlsx(db))).active
    finally:
        db.close()
    headers = [c.value for c in ws[2]]
    assert headers[:3] == ["매체 ID", "매체명", "구분명"]
    assert ws["A1"].value == "기본 정보"
    assert "디바이스 유형" not in headers and "등급 근거" not in headers
    assert headers[-2:] == ["대표 이미지 URL", "사진 수"]
    source_col = headers.index("고정/이동") + 1
    values = {ws.cell(row=r, column=source_col).value for r in range(3, ws.max_row + 1)}
    assert values <= {"고정", "이동", "", None}


def _filled_template(rows: list[dict], plans: list[dict]) -> bytes:
    """양식을 받아 칸 이름(* 뺀 것) 기준으로 매체 등록·상품 시트를 채운다."""
    wb = load_workbook(BytesIO(media_excel.media_template_xlsx()))
    for sheet, header_row, first_row, values in (
        ("매체 등록", 2, 3, rows),
        ("상품", 1, 2, plans),
    ):
        ws = wb[sheet]
        labels = [str(c.value or "").rstrip(" *") for c in ws[header_row]]
        for r, row in enumerate(values, start=first_row):
            for label, value in row.items():
                ws.cell(row=r, column=labels.index(label) + 1, value=value)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_layout():
    wb = load_workbook(BytesIO(media_excel.media_template_xlsx()))
    assert wb.sheetnames == ["매체 등록", "상품", "작성 안내"]
    ws = wb["매체 등록"]
    headers = [c.value for c in ws[2]]
    assert headers[0] == "연결 번호 *"
    assert "매체명 *" in headers and "구분명" in headers
    assert "매체 ID" not in headers and "원본 상세 ID" not in headers
    assert "상권(호칭)" not in headers and "지역 코드" not in headers
    assert "디바이스 유형" not in headers
    # 가격은 상품에서 계산하므로 매체 시트에 없다.
    assert "최소 광고비(원)" not in headers and "제작비 유무" not in headers
    # 위치는 도로명 주소·상세 주소만 적고 나머지는 올릴 때 채운다. 시/도·구/군은 이동 매체용.
    assert "도로명 주소 *" in headers and "상세 주소" in headers
    for gone in ("지번 주소", "도로명 주소(동 포함)", "건물명", "법정동", "위도", "경도"):
        assert gone not in headers
    assert "시/도 (이동 매체) *" in headers and "구/군 (이동 매체)" in headers
    plan_headers = [c.value for c in wb["상품"][1]]
    assert plan_headers[:3] == ["연결 번호 *", "상품명 *", "판매 방식 *"]
    assert ws.data_validations.dataValidation and wb["상품"].data_validations.dataValidation


def _plan(link, name, ad=None, prod=None, **kw):
    return {"연결 번호": link, "상품명": name, "판매 방식": kw.get("master", "개별"),
            "계약 기간": 1, "기간 단위": "개월", "광고비(원)": ad, "제작비(원)": prod}


def test_import_korean_template_with_plans():
    base = {
        "매체명": "엑셀등록 고정",
        "고정/이동": "고정",
        "판매 유형": "단품",
        "카테고리(대)": "전광판/빌보드",
        "카테고리(소)": "전광판",
        "매체 타입": "OOH",
        "설치 장소": "외부형",
        "도로명 주소": "서울 강남구 강남대로 396",
        "검색 태그": "강남, 대형|핫플",
        "인기 업종": "패션, 화장품",
    }
    moving = {**base, "매체명": "엑셀등록 이동", "고정/이동": "이동", "시/도 (이동 매체)": "서울특별시",
              "구/군 (이동 매체)": "강남구,서초구", "도로명 주소": None,
              "매체 타입": "DOOH"}
    content = _filled_template(
        [
            {**base, "연결 번호": 1},
            {**moving, "연결 번호": 2},
            {**base, "연결 번호": 3, "매체명": "필수 누락", "도로명 주소": None},  # 고정인데 주소 없음
            {**base, "연결 번호": 4, "매체명": "잘못된 값", "설치 장소": "옥상"},  # 선택지에 없음
            {**base, "연결 번호": 5, "매체명": "상품 없음"},                      # 상품 시트에 없음
            {**base, "연결 번호": 1, "매체명": "번호 겹침"},                      # 연결 번호 중복
        ],
        [
            _plan(1, "빌보드", "5,000,000", 300000),
            _plan(1, "랩핑", 3000000),
            _plan(2, "영상(20초)", 1000000, 50000),  # DOOH → 제작비 저장 안 함
            _plan(9, "주인 없음", 100),              # 연결 번호 9인 매체 없음
            {"연결 번호": 2, "상품명": "단위 없음", "판매 방식": "개별", "계약 기간": 1},  # 기간 단위 누락
        ],
    )
    db = SessionLocal()
    created: list[str] = []
    try:
        result = media_excel.import_media_xlsx(db, content)
        rows = db.query(Media).filter(Media.name.like("엑셀등록%")).all()
        created = [m.media_id for m in rows]
        by_name = {m.name: m for m in rows}

        assert result["inserted"] == 2 and result["failed"] == 4, result
        reasons = " ".join(e["reason"] for e in result["errors"])
        assert "도로명 주소" in reasons and "공간형 / 외부형" in reasons
        assert "상품이 없습니다" in reasons and "겹칩니다" in reasons
        assert "연결 번호 9인 매체가" in reasons and "기간 단위" in reasons

        fixed = by_name["엑셀등록 고정"]
        assert fixed.media_id.startswith("F")
        assert fixed.sales_type == "SINGLE" and fixed.exposure_type == "OUTSIDE"
        assert [p.product_display_name for p in fixed.plans] == ["빌보드", "랩핑"]
        assert fixed.min_advertisement_fee_krw == 3_000_000
        assert fixed.max_advertisement_fee_krw == 5_000_000
        assert fixed.min_production_fee_krw == 300_000 and fixed.any_production_fee_yn is True
        assert fixed.plan_count == 2
        assert fixed.list_labels == ["강남", "대형", "핫플"]
        assert fixed.popular_type == "패션,화장품"

        # 기기·면 수량은 비우면 1
        assert {(p.default_device_quantity, p.default_surface_quantity) for p in fixed.plans} == {(1, 1)}

        mv = by_name["엑셀등록 이동"]
        assert mv.media_id.startswith("M") and mv.city == "서울특별시"
        assert len(mv.plans) == 1 and mv.plans[0].production_fee is None
        assert mv.any_production_fee_yn is False
    finally:
        if created:
            db.query(MediaPlan).filter(MediaPlan.media_id.in_(created)).delete(synchronize_session=False)
            db.query(Media).filter(Media.media_id.in_(created)).delete(synchronize_session=False)
            db.commit()
        db.close()


def test_export_has_plan_sheet():
    db = SessionLocal()
    try:
        wb = load_workbook(BytesIO(media_excel.export_media_xlsx(db)))
        plan_count = db.query(MediaPlan).count()
    finally:
        db.close()
    ws = wb["상품"]
    assert [c.value for c in ws[1]][:4] == ["매체 ID", "매체명", "상품 번호", "상품명"]
    assert ws.max_row - 1 == plan_count


def test_import_downloaded_file_skips_existing_ids():
    db = SessionLocal()
    try:
        before = db.query(Media).count()
        result = media_excel.import_media_xlsx(db, media_excel.export_media_xlsx(db))
        assert result["inserted"] == 0 and result["skipped"] == before
        assert result["errors"] == []  # 건너뛴 매체의 상품 줄은 오류로 잡지 않는다
        assert db.query(Media).count() == before
    finally:
        db.close()


def test_old_header_labels_still_read():
    """칸 이름이 바뀌기 전에 받아 둔 엑셀도 읽는다."""
    assert media_excel._header_key_map()["OOH 유형"][0] == "ooh_type"


def test_old_exposure_values_still_read():
    assert media_excel._parse_cell("exposure_type", {"INSIDE": "공간형", "OUTSIDE": "외부형"}, "실내") == "INSIDE"
    assert media_excel._parse_cell("exposure_type", {"INSIDE": "공간형", "OUTSIDE": "외부형"}, "외부형") == "OUTSIDE"


def test_old_shape_values_still_read():
    shapes = dict(next(c for _, cols in media_excel.SECTIONS for c in cols if c[0] == "media_shape")[2])
    assert shapes["CURVED_SHAPE"] == "커브형" and shapes["DIFFERENT_SHAPE"] == "변형"
    assert media_excel._parse_cell("media_shape", shapes, "곡면형") == "CURVED_SHAPE"
    assert media_excel._parse_cell("media_shape", shapes, "변형") == "DIFFERENT_SHAPE"


def test_formats_and_time_cells():
    from datetime import time

    assert media_excel._parse_cell("material_formats", "formats", "영상(MP4), 스틸컷(JPG·PNG), MP4") == ["MP4", "IMAGE"]
    assert media_excel._parse_cell("operation_start_time", "time", time(6, 0)) == "06:00"
    assert media_excel._parse_cell("operation_end_time", "time", "24:00") == "24:00"
    assert media_excel._cell(["MP4", "IMAGE"], "formats") == "영상(MP4), 스틸컷(JPG·PNG)"


def test_import_fills_address_from_road_address(fake_kakao_address):
    base = {
        "매체명": "엑셀주소 자동",
        "고정/이동": "고정",
        "판매 유형": "단품",
        "카테고리(대)": "전광판/빌보드",
        "카테고리(소)": "전광판",
        "매체 타입": "OOH",
        "설치 장소": "외부형",
        "도로명 주소": "서울 강남구 강남대로 396",
        "상세 주소": "2층 로비",
    }
    content = _filled_template(
        [
            {**base, "연결 번호": 1},
            # 같은 주소 — 검색은 한 번만, 직접 적은 시/도는 그대로
            {**base, "연결 번호": 2, "매체명": "엑셀주소 직접", "시/도 (이동 매체)": "서울특별시"},
            {**base, "연결 번호": 3, "매체명": "엑셀주소 없음", "도로명 주소": "없는 주소 1"},
        ],
        [_plan(n, "빌보드", 100) for n in (1, 2, 3)],
    )
    db = SessionLocal()
    created: list[str] = []
    try:
        result = media_excel.import_media_xlsx(db, content)
        rows = db.query(Media).filter(Media.name.like("엑셀주소%")).all()
        created = [m.media_id for m in rows]
        by_name = {m.name: m for m in rows}

        assert result["inserted"] == 2 and result["failed"] == 1, result
        assert "주소를 찾지 못했습니다: 없는 주소 1" in result["errors"][0]["reason"]
        assert fake_kakao_address.count("서울 강남구 강남대로 396") == 1

        auto = by_name["엑셀주소 자동"]
        assert auto.accurate_address == "서울 강남구 강남대로 396 (역삼동)"
        assert auto.full_address_jibun == "서울 강남구 역삼동 858"
        assert auto.building_name == "강남역 빌딩" and auto.legal_dong == "역삼동"
        assert (auto.city, auto.district) == ("서울", "강남구")
        assert float(auto.latitude) == 37.4979 and float(auto.longitude) == 127.0276
        assert auto.address_detail == "2층 로비"

        manual = by_name["엑셀주소 직접"]
        assert manual.city == "서울특별시" and manual.legal_dong == "역삼동"
        assert float(manual.latitude) == 37.4979
    finally:
        if created:
            db.query(MediaPlan).filter(MediaPlan.media_id.in_(created)).delete(synchronize_session=False)
            db.query(Media).filter(Media.media_id.in_(created)).delete(synchronize_session=False)
            db.commit()
        db.close()


def test_kakao_lookup_maps_response(monkeypatch):
    class Res:
        status_code = 200

        @staticmethod
        def json():
            return {"documents": [{
                "x": "127.0276", "y": "37.4979",
                "address": {"address_name": "서울 강남구 역삼동 858", "region_1depth_name": "서울",
                            "region_2depth_name": "강남구", "region_3depth_name": "역삼동"},
                "road_address": {"address_name": "서울 강남구 강남대로 396", "building_name": "",
                                 "region_1depth_name": "서울", "region_2depth_name": "강남구",
                                 "region_3depth_name": "역삼동"},
            }]}

    class Client:
        def get(self, url, params, headers, timeout):
            assert params == {"query": "서울 강남구 강남대로 396"}
            assert headers["Authorization"] == "KakaoAK test-key"
            return Res()

    monkeypatch.undo()  # autouse 가짜 lookup 을 걷어 내고 진짜 함수를 본다
    monkeypatch.setattr(kakao_address, "_api_key", lambda: "test-key")
    found = kakao_address.lookup("서울 강남구 강남대로 396", client=Client())
    assert found == {
        "accurate_address": "서울 강남구 강남대로 396 (역삼동)",
        "full_address_jibun": "서울 강남구 역삼동 858",
        "legal_dong": "역삼동",
        "city": "서울",
        "district": "강남구",
        "latitude": 37.4979,
        "longitude": 127.0276,
    }
