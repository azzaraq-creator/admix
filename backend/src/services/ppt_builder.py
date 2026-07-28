"""제안서 PPT 생성 — python-pptx 로 슬라이드 레이아웃을 코드로 구성하고 데이터를 채운다.

프론트 React 슬라이드(components/proposals/*Template.tsx)의 1920x1080 디자인을
근사 재현한다. 슬라이드 순서: 표지 → 서머리(매체 5개씩) → 매체별 → Thanks.
"""
from __future__ import annotations

import io
import os
from typing import Optional

import httpx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from src.config import get_settings

# 1920x1080 px 디자인 → 슬라이드 16:9 (12192000 x 6858000 EMU)
EMU_PER_PX = 6350  # 1920 * 6350 = 12192000
PT_PER_PX = 0.5  # 1920px = 960pt (13.333in)
SLIDE_W = Emu(1920 * EMU_PER_PX)
SLIDE_H = Emu(1080 * EMU_PER_PX)

FONT = "Noto Sans KR"
TEAL = RGBColor(0x00, 0xAA, 0xA4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_TEXT = RGBColor(0x54, 0x54, 0x54)
DARK = RGBColor(0x2A, 0x2A, 0x2A)
HEADER_BG = RGBColor(0xF6, 0xF6, 0xF6)
BORDER = RGBColor(0xE4, 0xE5, 0xEE)
PLACEHOLDER_BG = RGBColor(0xF6, 0xF6, 0xF6)
PLACEHOLDER_TX = RGBColor(0xA0, 0xA0, 0xA0)

ROWS_PER_PAGE = 5
EMPTY = "-"

# 표지·Thanks 풀블리드 배경 이미지 (1920×1080)
COVER_BG = os.path.join(os.path.dirname(__file__), "..", "assets", "cover-bg.png")

SUMMARY_COLS = [
    ("NO", 84),
    ("구분", 168),
    ("지역", 168),
    ("매체명", 219),
    ("상품명", 219),
    ("수량", 126),
    ("광고비", 218),
    ("제작비", 218),
    ("합계", 218),
    ("시작일/종료일", 202),
]
MEDIA_COLS = [
    ("상품명", 290),
    ("노출회수", 219),
    ("판매수량", 219),
    ("규격", 219),
    ("운영시간", 219),
    ("기간", 219),
    ("광고비", 218),
    ("제작비용", 218),
]


def _px(v: float) -> Emu:
    return Emu(int(v * EMU_PER_PX))


def _pt(v: float):
    return Pt(v * PT_PER_PX)


def _won(value) -> str:
    try:
        return f"{int(value):,}원"
    except (TypeError, ValueError):
        return EMPTY


def _num(value) -> str:
    try:
        return f"{int(value):,}"
    except (TypeError, ValueError):
        return EMPTY


def _year_date(iso: Optional[str]) -> tuple[str, str]:
    from datetime import datetime, timezone

    if iso:
        try:
            d = datetime.fromisoformat(iso.replace("Z", "+00:00"))
        except ValueError:
            d = datetime.now(timezone.utc)
    else:
        d = datetime.now(timezone.utc)
    return str(d.year), f"{d.year}.{d.month:02d}.{d.day:02d}"


def _fetch_image(url: Optional[str]) -> Optional[bytes]:
    if not url:
        return None
    try:
        resp = httpx.get(url, timeout=10, follow_redirects=True)
        resp.raise_for_status()
        return resp.content
    except Exception:  # noqa: BLE001 — 이미지 실패는 graceful
        return None


def _bg(slide, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _bg_pattern(slide) -> None:
    """표지·Thanks 공용 배경 — 1920×1080 배경 이미지를 슬라이드 전체에 풀블리드."""
    if os.path.exists(COVER_BG):
        slide.shapes.add_picture(COVER_BG, 0, 0, width=SLIDE_W, height=SLIDE_H)
    else:
        _bg(slide, RGBColor(0x5A, 0x5A, 0x5A))


def _rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size,
    color,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = _pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    return box


def _placeholder_box(slide, left, top, width, height, label: str) -> None:
    box = _rect(slide, left, top, width, height, PLACEHOLDER_BG)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = _pt(28)
    run.font.name = FONT
    run.font.color.rgb = PLACEHOLDER_TX


def _image_or_placeholder(slide, left, top, width, height, url, label):
    data = _fetch_image(url)
    if data is None:
        _placeholder_box(slide, left, top, width, height, label)
        return
    try:
        slide.shapes.add_picture(io.BytesIO(data), left, top, width=width, height=height)
    except Exception:  # noqa: BLE001
        _placeholder_box(slide, left, top, width, height, label)


def _set_cell(cell, text, *, size=18, bold=False, color=GRAY_TEXT, bg=None):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = _px(8)
    cell.margin_right = _px(8)
    cell.margin_top = 0
    cell.margin_bottom = 0
    if bg is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(text)
    run.font.size = _pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color


def _new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


# ===== 슬라이드 빌더 =====


def build_cover(prs, detail: dict) -> None:
    slide = _new_slide(prs)
    _bg_pattern(slide)
    year, date = _year_date(detail.get("updated_at"))
    _rect(slide, _px(80), _px(80), _px(740), _px(24), WHITE)
    _text(slide, _px(80), _px(140), _px(1200), _px(320), "ADMIX\n제안서",
          size=120, color=WHITE, bold=True)
    _text(slide, _px(80), _px(800), _px(1700), _px(160), f"광고 제안서_{year}",
          size=120, color=WHITE, bold=True)
    _text(slide, _px(80), _px(980), _px(1700), _px(70), date,
          size=48, color=RGBColor(0xCC, 0xCC, 0xCC))


def build_thanks(prs) -> None:
    slide = _new_slide(prs)
    _bg_pattern(slide)
    _text(slide, 0, _px(460), SLIDE_W, _px(160), "THANK YOU",
          size=120, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)


def build_summary_pages(prs, detail: dict) -> None:
    items = detail.get("items", [])
    pages = [items[i:i + ROWS_PER_PAGE] for i in range(0, len(items), ROWS_PER_PAGE)]
    if not pages:
        pages = [[]]

    ad_total = sum((it.get("price") or 0) * (it.get("quantity") or 1) for it in items)
    prod_total = sum(
        (it.get("production_fee") or 0) * (it.get("quantity") or 1) for it in items
    )
    regions = ", ".join(
        dict.fromkeys(it.get("region") for it in items if it.get("region"))
    ) or EMPTY
    starts = sorted(it["start_date"] for it in items if it.get("start_date"))
    ends = sorted((it["end_date"] for it in items if it.get("end_date")), reverse=True)
    period = f"{starts[0]} ~ {ends[0]}" if starts and ends else EMPTY

    for page_idx, rows in enumerate(pages):
        slide = _new_slide(prs)
        # 청록 헤더
        _rect(slide, 0, 0, SLIDE_W, _px(340), TEAL)
        _text(slide, _px(80), _px(40), _px(500), _px(110), "Summary",
              size=80, color=WHITE, bold=True)
        _text(slide, _px(80), _px(200), _px(400), _px(40), "캠페인 기간",
              size=24, color=WHITE)
        _text(slide, _px(80), _px(245), _px(560), _px(50), period,
              size=28, color=WHITE, bold=True)

        stats = [
            ("집행 매체", str(len(items))),
            ("캠페인 지역", regions),
            ("광고비 합계", _won(ad_total)),
            ("제작비 합계", _won(prod_total)),
            ("전체 금액 합계(GROSS) *VAT 별도", _won(detail.get("total_amount"))),
        ]
        # 헤더 우측 통계 2열 + 마지막 전체금액
        positions = [
            (680, 60), (680, 190), (1080, 60), (1080, 190), (1480, 60),
        ]
        for (label, value), (lx, ty) in zip(stats, positions):
            _text(slide, _px(lx), _px(ty), _px(380), _px(40), label,
                  size=22, color=WHITE)
            _text(slide, _px(lx), _px(ty + 45), _px(420), _px(60), value,
                  size=28 if label.startswith("전체") else 26, color=WHITE, bold=True)

        # 표
        n_rows = 1 + ROWS_PER_PAGE
        table_left, table_top = _px(40), _px(380)
        table_w = _px(sum(w for _, w in SUMMARY_COLS))
        table_h = _px(640)
        graphic = slide.shapes.add_table(n_rows, len(SUMMARY_COLS), table_left,
                                         table_top, table_w, table_h)
        table = graphic.table
        table.first_row = False
        for ci, (label, w) in enumerate(SUMMARY_COLS):
            table.columns[ci].width = _px(w)
            _set_cell(table.cell(0, ci), label, bold=True, bg=HEADER_BG)
        for ri in range(ROWS_PER_PAGE):
            it = rows[ri] if ri < len(rows) else None
            no = page_idx * ROWS_PER_PAGE + ri + 1 if it else ""
            q = (it.get("quantity") or 1) if it else 1
            vals = (
                [
                    no,
                    it.get("category") or EMPTY,
                    it.get("region") or EMPTY,
                    it.get("name") or EMPTY,
                    it.get("product") or EMPTY,
                    _num(q),
                    _num((it.get("price") or 0) * q),
                    _num((it.get("production_fee") or 0) * q),
                    _num(((it.get("price") or 0) + (it.get("production_fee") or 0)) * q),
                    f"{it.get('start_date') or EMPTY}\n{it.get('end_date') or EMPTY}",
                ]
                if it
                else [""] * len(SUMMARY_COLS)
            )
            for ci, v in enumerate(vals):
                _set_cell(table.cell(ri + 1, ci), v)


def build_media(prs, item: dict) -> None:
    slide = _new_slide(prs)
    _bg(slide, WHITE)

    # 좌측 매체 정보
    _text(slide, _px(38), _px(40), _px(480), _px(80), item.get("name") or EMPTY,
          size=32, color=DARK, bold=True)
    meta = " · ".join(
        v for v in (item.get("address"), item.get("ooh_type"), item.get("category")) if v
    )
    if meta:
        _text(slide, _px(38), _px(120), _px(480), _px(60), meta,
              size=16, color=RGBColor(0x73, 0x75, 0x86))
    if item.get("description"):
        _text(slide, _px(38), _px(190), _px(480), _px(300), item["description"],
              size=16, color=DARK)

    # 썸네일 이미지
    _image_or_placeholder(slide, _px(548), _px(40), _px(822), _px(500),
                          item.get("thumbnail_url"), "이미지 없음")

    # 지도
    settings = get_settings()
    lat, lng = item.get("latitude"), item.get("longitude")
    map_url = None
    if lat is not None and lng is not None and settings.geoapify_api_key:
        map_url = (
            "https://maps.geoapify.com/v1/staticmap"
            "?style=osm-bright&width=500&height=500&scaleFactor=2"
            f"&center=lonlat:{lng},{lat}&zoom=16"
            f"&marker=lonlat:{lng},{lat};color:%23ff0000;size:medium"
            f"&apiKey={settings.geoapify_api_key}"
        )
    _image_or_placeholder(slide, _px(1382), _px(40), _px(500), _px(500),
                          map_url, "위치 정보 없음")

    # 선택 plan
    plans = item.get("plans") or []
    plan = next((p for p in plans if p.get("plan_no") == item.get("selected_plan_no")),
                plans[0] if plans else None)
    product = (
        (plan.get("product_display_name") or plan.get("product_name")) if plan else None
    ) or item.get("product") or EMPTY
    op = (
        f"{plan['operation_start_time']}~{plan['operation_end_time']}"
        if plan and plan.get("operation_start_time") and plan.get("operation_end_time")
        else EMPTY
    )
    qty = (
        f"{item.get('device_quantity') or EMPTY}기 {item.get('surface_quantity') or EMPTY}면"
    )

    # 하단 표
    table_left, table_top = _px(38), _px(580)
    table_w = _px(sum(w for _, w in MEDIA_COLS))
    graphic = slide.shapes.add_table(2, len(MEDIA_COLS), table_left, table_top,
                                     table_w, _px(180))
    table = graphic.table
    table.first_row = False
    for ci, (label, w) in enumerate(MEDIA_COLS):
        table.columns[ci].width = _px(w)
        _set_cell(table.cell(0, ci), label, bold=True, bg=HEADER_BG)
    vals = [
        product,
        EMPTY,
        qty,
        item.get("spec") or EMPTY,
        op,
        EMPTY,
        _num(item.get("price")),
        _num(plan.get("production_fee") if plan else None),
    ]
    for ci, v in enumerate(vals):
        _set_cell(table.cell(1, ci), v)


def generate_proposal_ppt(detail: dict, out_path: str) -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    build_cover(prs, detail)
    build_summary_pages(prs, detail)
    for item in detail.get("items", []):
        build_media(prs, item)
    build_thanks(prs)
    prs.save(out_path)
    return len(prs.slides._sldIdLst)
