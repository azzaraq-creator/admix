"""매체 목록 조회 — admin/media 페이지용 평면 매핑.

media 단일값 + 대표 플랜(plan_no=1)의 상품표시명을 합쳐 한 행으로 만든다.
"""
from __future__ import annotations

import json
import math
import os
import uuid
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import (
    BigInteger,
    Boolean,
    DateTime,
    Integer,
    Numeric,
    and_,
    func,
    inspect as sa_inspect,
    text,
)
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.orm import Session, selectinload

from src.config import get_settings
from src.models.media_image import MediaImage
from src.models.media_master import Media
from src.models.media_plan import MediaPlan
from src.services.graph.settings import SANGWON_MAX_DISTANCE_M, SANGWON_QUARTER

# admin 매체 상세/등록 폼 — 자동 관리 컬럼(수정 대상 아님)
_MEDIA_AUTO_COLS = {"created_at", "updated_at"}
_MEDIA_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
_MEDIA_IMAGE_MAX_BYTES = 10 * 1024 * 1024  # 10MB

_SALE_TYPE = {"SINGLE": "단품", "GROUP": "묶음"}
_EXPOSURE_TYPE = {"INSIDE": "실내", "OUTSIDE": "실외"}
_DURATION_TYPE = {
    "YEARS": "년",
    "MONTHS": "개월",
    "WEEKS": "주",
    "DAYS": "일",
    "HOURS": "시간",
}


def _media_badge(m: Media) -> str | None:
    if m.is_popular_yn:
        return "popular"
    if m.is_newly_built_yn:
        return "new"
    return None


def _fmt_fee(v: int | None) -> str:
    return f"{v:,}" if v is not None else "-"


def _fmt_date(dt: datetime | None) -> str:
    return dt.date().isoformat() if dt is not None else "-"


def _media_card(m: Media) -> dict:
    name = " ".join(p for p in [(m.name or "").strip(), (m.second_name or "").strip()] if p)
    # 카드 이미지: 썸네일 우선 → media_image(sort_order) 순, 중복 제거 후 최대 3장.
    images: list[str] = []
    for url in [m.thumbnail_url, *(img.image_url for img in m.images)]:
        if url and url not in images:
            images.append(url)
    return dict(
        id=m.media_id,
        name=name or "-",
        minAdvertisementFeeKrw=m.min_advertisement_fee_krw,
        thumbnailUrl=m.thumbnail_url,
        images=images[:3],
        badge=_media_badge(m),
        lat=float(m.latitude) if m.latitude is not None else None,
        lng=float(m.longitude) if m.longitude is not None else None,
    )


def list_moving_media(
    db: Session,
    *,
    categories: list[str] | None = None,
    ooh_types: list[str] | None = None,
    exposure_types: list[str] | None = None,
    media_shapes: list[str] | None = None,
    product_master_types: list[str] | None = None,
    price_min: int | None = None,
    price_max: int | None = None,
) -> list[dict]:
    base = _media_base_query(
        db,
        media_source="MOVING",
        categories=categories,
        ooh_types=ooh_types,
        exposure_types=exposure_types,
        media_shapes=media_shapes,
        product_master_types=product_master_types,
        price_min=price_min,
        price_max=price_max,
    )
    rows = base.options(selectinload(Media.images)).order_by(Media.media_id).all()
    return [_media_card(m) for m in rows]


def _media_base_query(
    db: Session,
    *,
    media_source: str = "FIXED",
    categories: list[str] | None,
    ooh_types: list[str] | None,
    exposure_types: list[str] | None,
    media_shapes: list[str] | None,
    product_master_types: list[str] | None,
    price_min: int | None,
    price_max: int | None,
    ne_lat: float | None = None,
    sw_lat: float | None = None,
    ne_lng: float | None = None,
    sw_lng: float | None = None,
):
    """매체 공통 필터 쿼리. 리스트·지도클러스터가 동일 조건을 공유한다."""
    base = db.query(Media).filter(Media.media_source == media_source)
    if categories:
        base = base.filter(Media.category_large.in_(categories))
    if ooh_types:
        base = base.filter(Media.ooh_type.in_(ooh_types))
    if exposure_types:
        base = base.filter(Media.exposure_type.in_(exposure_types))
    if media_shapes:
        base = base.filter(Media.media_shape.in_(media_shapes))
    if price_min is not None:
        base = base.filter(Media.min_advertisement_fee_krw >= price_min)
    if price_max is not None:
        base = base.filter(Media.min_advertisement_fee_krw <= price_max)
    if product_master_types:
        # 매체판매유형 = 해당 매체 플랜의 product_master_type 집합 기준(조인 IN)
        plan_media_ids = db.query(MediaPlan.media_id).filter(
            MediaPlan.product_master_type.in_(product_master_types)
        )
        base = base.filter(Media.media_id.in_(plan_media_ids))
    if None not in (ne_lat, sw_lat, ne_lng, sw_lng):
        # 지도 화면 영역(bounding box) 종속 — 리스트=지도 영역.
        base = base.filter(
            Media.latitude.isnot(None),
            Media.longitude.isnot(None),
            Media.latitude >= sw_lat,
            Media.latitude <= ne_lat,
            Media.longitude >= sw_lng,
            Media.longitude <= ne_lng,
        )
    return base


def list_fixed_media(
    db: Session,
    *,
    limit: int,
    offset: int,
    categories: list[str] | None = None,
    ooh_types: list[str] | None = None,
    exposure_types: list[str] | None = None,
    media_shapes: list[str] | None = None,
    product_master_types: list[str] | None = None,
    price_min: int | None = None,
    price_max: int | None = None,
    ne_lat: float | None = None,
    sw_lat: float | None = None,
    ne_lng: float | None = None,
    sw_lng: float | None = None,
) -> tuple[int, list[dict]]:
    base = _media_base_query(
        db,
        categories=categories,
        ooh_types=ooh_types,
        exposure_types=exposure_types,
        media_shapes=media_shapes,
        product_master_types=product_master_types,
        price_min=price_min,
        price_max=price_max,
        ne_lat=ne_lat,
        sw_lat=sw_lat,
        ne_lng=ne_lng,
        sw_lng=sw_lng,
    )
    total = base.count()
    rows = (
        base.options(selectinload(Media.images))
        .order_by(Media.media_id)
        .limit(limit)
        .offset(offset)
        .all()
    )
    return total, [_media_card(m) for m in rows]


# 지도 마커 클러스터링 — zoom_level 기반 그리드 셀 크기(도 단위).
# kakao 지도 레벨은 1=최대확대 … 14=최대축소. 축소(level↑)일수록 셀이 커져 더 많이 묶인다.
_CLUSTER_CELL_DEG_BASE = 0.0006
_CLUSTER_MIN_LEVEL = 1
# 이 레벨 이하(= 더 확대)에선 클러스터를 만들지 않고 개별 마커(핀)만 표시.
# 1 = 최대 확대(레벨1)일 때만 개별 핀. 레벨 2+에선 그리드 클러스터링이 줌에 따라 변한다.
_CLUSTER_DECLUSTER_LEVEL = 1


def _cluster_cell_deg(zoom_level: int) -> float:
    step = max(0, zoom_level - _CLUSTER_MIN_LEVEL)
    return _CLUSTER_CELL_DEG_BASE * (2 ** step)


def _images_by_media(db: Session, media_ids: list[str]) -> dict[str, list[str]]:
    """media_id → media_image url 목록(sort_order). 마커 카드 이미지용 일괄 조회."""
    if not media_ids:
        return {}
    rows = (
        db.query(MediaImage.media_id, MediaImage.image_url)
        .filter(MediaImage.media_id.in_(media_ids))
        .order_by(MediaImage.media_id, MediaImage.sort_order)
        .all()
    )
    out: dict[str, list[str]] = {}
    for mid, url in rows:
        out.setdefault(mid, []).append(url)
    return out


def _marker_from_row(r, lat: float, lng: float, img_map: dict) -> dict:
    name = " ".join(
        p for p in [(r.name or "").strip(), (r.second_name or "").strip()] if p
    )
    if r.is_popular_yn:
        badge = "popular"
    elif r.is_newly_built_yn:
        badge = "new"
    else:
        badge = None
    # 카드 이미지: 썸네일 우선 → media_image(sort_order), 중복 제거 최대 3장 — 검색 리스트(_media_card)와 동일.
    images: list[str] = []
    for url in [r.thumbnail_url, *img_map.get(r.media_id, [])]:
        if url and url not in images:
            images.append(url)
    images = images[:3]
    return dict(
        id=r.media_id,
        lat=lat,
        lng=lng,
        name=name or "-",
        categoryLarge=r.category_large,
        minAdvertisementFeeKrw=r.min_advertisement_fee_krw,
        thumbnailUrl=images[0] if images else None,
        images=images,
        badge=badge,
    )


def list_fixed_clusters(
    db: Session,
    *,
    zoom_level: int,
    ne_lat: float,
    sw_lat: float,
    ne_lng: float,
    sw_lng: float,
    categories: list[str] | None = None,
    ooh_types: list[str] | None = None,
    exposure_types: list[str] | None = None,
    media_shapes: list[str] | None = None,
    product_master_types: list[str] | None = None,
    price_min: int | None = None,
    price_max: int | None = None,
) -> dict:
    """지도 화면(bbox) 안 FIXED 매체를 zoom_level 그리드로 묶어 클러스터/마커로 반환."""
    base = _media_base_query(
        db,
        categories=categories,
        ooh_types=ooh_types,
        exposure_types=exposure_types,
        media_shapes=media_shapes,
        product_master_types=product_master_types,
        price_min=price_min,
        price_max=price_max,
        ne_lat=ne_lat,
        sw_lat=sw_lat,
        ne_lng=ne_lng,
        sw_lng=sw_lng,
    )
    rows = base.with_entities(
        Media.media_id,
        Media.name,
        Media.second_name,
        Media.category_large,
        Media.latitude,
        Media.longitude,
        Media.min_advertisement_fee_krw,
        Media.thumbnail_url,
        Media.is_popular_yn,
        Media.is_newly_built_yn,
    ).all()

    points = [
        (r, float(r.latitude), float(r.longitude))
        for r in rows
        if r.latitude is not None and r.longitude is not None
    ]

    # 마커로 나갈 행 결정(확대 임계치 이하: 전부 개별 핀 / 그 외: 1개짜리 버킷만).
    clusters: list[dict] = []
    if zoom_level <= _CLUSTER_DECLUSTER_LEVEL:
        marker_points = points
    else:
        cell = _cluster_cell_deg(zoom_level)
        buckets: dict[tuple[int, int], list] = {}
        for r, lat, lng in points:
            key = (math.floor(lat / cell), math.floor(lng / cell))
            buckets.setdefault(key, []).append((r, lat, lng))
        marker_points = []
        for members in buckets.values():
            if len(members) == 1:
                marker_points.append(members[0])
            else:
                n = len(members)
                clusters.append(
                    dict(
                        lat=sum(m[1] for m in members) / n,
                        lng=sum(m[2] for m in members) / n,
                        count=n,
                    )
                )

    # 마커 매체 이미지 일괄 조회 → 검색 리스트와 동일한 카드 이미지(최대 3장).
    img_map = _images_by_media(db, [r.media_id for r, _, _ in marker_points])
    markers = [
        _marker_from_row(r, lat, lng, img_map) for r, lat, lng in marker_points
    ]
    return dict(clusters=clusters, markers=markers)


def get_media_filter_options(db: Session, media_source: str = "FIXED") -> dict:
    """매체검색 필터 옵션 — 매체 기준 distinct 값 + 가격(최소광고비) 범위."""

    def _distinct(col) -> list[str]:
        rows = (
            db.query(col)
            .filter(Media.media_source == media_source, col.isnot(None))
            .distinct()
            .order_by(col)
            .all()
        )
        return [r[0] for r in rows]

    pmt_rows = (
        db.query(MediaPlan.product_master_type)
        .join(Media, MediaPlan.media_id == Media.media_id)
        .filter(
            Media.media_source == media_source,
            MediaPlan.product_master_type.isnot(None),
        )
        .distinct()
        .order_by(MediaPlan.product_master_type)
        .all()
    )
    price = (
        db.query(
            func.min(Media.min_advertisement_fee_krw),
            func.max(Media.min_advertisement_fee_krw),
        )
        .filter(Media.media_source == media_source)
        .first()
    )
    price_min = price[0] if price else None
    price_max = price[1] if price else None
    return dict(
        categories=_distinct(Media.category_large),
        ooh_types=_distinct(Media.ooh_type),
        exposure_types=_distinct(Media.exposure_type),
        media_shapes=_distinct(Media.media_shape),
        product_master_types=[r[0] for r in pmt_rows],
        price_min=price_min,
        price_max=price_max,
        price_histogram=_price_histogram(db, price_min, price_max, media_source),
    )


_PRICE_HISTOGRAM_BUCKETS = 24


def _price_histogram(
    db: Session,
    price_min: int | None,
    price_max: int | None,
    media_source: str = "FIXED",
) -> list[int]:
    """min~max 가격 구간을 균등 버킷으로 나눠 버킷별 매체 수를 센다."""
    if price_min is None or price_max is None or price_max <= price_min:
        return []
    fees = (
        db.query(Media.min_advertisement_fee_krw)
        .filter(
            Media.media_source == media_source,
            Media.min_advertisement_fee_krw.isnot(None),
        )
        .all()
    )
    n = _PRICE_HISTOGRAM_BUCKETS
    span = price_max - price_min
    counts = [0] * n
    for (fee,) in fees:
        idx = int((fee - price_min) * n / span)
        if idx >= n:
            idx = n - 1
        elif idx < 0:
            idx = 0
        counts[idx] += 1
    return counts


def _plan_subtitle(p: MediaPlan) -> str | None:
    parts: list[str] = []
    device_qty = p.active_device_quantity or p.default_device_quantity
    surface_qty = p.active_surface_quantity or p.default_surface_quantity
    if device_qty and surface_qty:
        parts.append(f"{device_qty}기 {surface_qty}면")
    if p.exposure_duration_seconds:
        parts.append(f"{p.exposure_duration_seconds}초")
    if p.contractual_duration and p.contractual_duration_type:
        unit = _DURATION_TYPE.get(p.contractual_duration_type, p.contractual_duration_type)
        parts.append(f"{p.contractual_duration}{unit}")
    return " / ".join(parts) or None


def _media_population(db: Session, source_detail_id: int | None) -> dict | None:
    """media.source_detail_id → ad_media.media_id → 인접 상권(sangwon) 유동인구.

    rerank_sangwon 과 동일 분기/거리 기준. 거리 초과·미매칭이면 None(섹션 숨김).
    """
    if source_detail_id is None:
        return None
    row = db.execute(
        text(
            """
            SELECT sp.sangwon_name, sp.total_foot_traffic,
                   sp.male_foot, sp.female_foot,
                   sp.age_10_foot, sp.age_20_foot, sp.age_30_foot,
                   sp.age_40_foot, sp.age_50_foot, sp.age_60_foot
            FROM ad_media am
            JOIN sangwon_population sp
              ON sp.sangwon_code = am.sangwon_code
             AND sp.quarter_code = :quarter
            WHERE am.media_id = :sid
              AND am.sangwon_distance_m <= :max_dist
            LIMIT 1
            """
        ),
        {"quarter": SANGWON_QUARTER, "sid": source_detail_id, "max_dist": SANGWON_MAX_DISTANCE_M},
    ).first()
    if row is None or not row.total_foot_traffic:
        return None

    total = float(row.total_foot_traffic)
    male = float(row.male_foot or 0)
    female = float(row.female_foot or 0)
    gender_base = male + female
    male_pct = round(male / gender_base * 100) if gender_base else 0
    age_buckets = [
        ("10", row.age_10_foot, "under"),
        ("20대", row.age_20_foot, None),
        ("30대", row.age_30_foot, None),
        ("40대", row.age_40_foot, None),
        ("50대", row.age_50_foot, None),
        ("60", row.age_60_foot, "over"),
    ]
    age_ratios = [
        dict(label=label, value=round((val or 0) / total * 100, 1), bound=bound)
        for label, val, bound in age_buckets
    ]
    return dict(
        sangwonName=row.sangwon_name,
        monthlyFootTraffic=int(row.total_foot_traffic),
        malePct=male_pct,
        femalePct=(100 - male_pct) if gender_base else 0,
        ageRatios=age_ratios,
    )


def get_media_detail(db: Session, media_id: str) -> dict | None:
    m = db.query(Media).filter(Media.media_id == media_id).first()
    if m is None:
        return None

    name = " ".join(p for p in [(m.name or "").strip(), (m.second_name or "").strip()] if p)

    badge = _media_badge(m)

    features: list[dict] = []

    def add(label: str, value: str | None) -> None:
        if value not in (None, "", "-"):
            features.append(dict(label=label, value=str(value)))

    add("매체 카테고리", m.category_small)
    add("타입", m.ooh_type)
    add("노출 종류", _EXPOSURE_TYPE.get(m.exposure_type, m.exposure_type))
    add("판매 형태", _SALE_TYPE.get(m.sales_type, m.sales_type))
    add("고정/이동", "이동" if m.media_source == "MOVING" else "고정")
    if m.device_quantity:
        add("기기 수량", f"{m.device_quantity}기")
    if m.lead_time_bizdays:
        add("리드타임", f"{m.lead_time_bizdays}영업일")

    plans = [
        dict(
            planNo=p.plan_no,
            title=p.product_display_name or p.product_name or "-",
            subtitle=_plan_subtitle(p),
        )
        for p in m.plans
        if p.product_master_type == "PM_INDIVIDUAL"
    ]

    return dict(
        id=m.media_id,
        name=name or "-",
        badge=badge,
        minAdvertisementFeeKrw=m.min_advertisement_fee_krw,
        maxAdvertisementFeeKrw=m.max_advertisement_fee_krw,
        description=m.description,
        address=(m.accurate_address or m.address or m.full_address_jibun or None),
        thumbnailUrl=m.thumbnail_url,
        imageUrls=[img.image_url for img in m.images],
        sizeText=(m.media_shape_summary or None),
        features=features,
        plans=plans,
        population=_media_population(db, m.source_detail_id),
    )


def list_media(db: Session) -> list[dict]:
    rows = (
        db.query(Media, MediaPlan)
        .outerjoin(
            MediaPlan,
            and_(MediaPlan.media_id == Media.media_id, MediaPlan.plan_no == 1),
        )
        .order_by(Media.media_id)
        .all()
    )
    items: list[dict] = []
    for m, plan in rows:
        product = "-"
        if plan is not None:
            product = plan.product_display_name or plan.product_name or "-"
        name = " ".join(p for p in [(m.name or "").strip(), (m.second_name or "").strip()] if p)
        items.append(
            dict(
                no=m.media_id,
                mediaType="이동" if m.media_source == "MOVING" else "고정",
                name=name or "-",
                region=m.loc_label or "-",
                category=m.category_small or "-",
                product=product,
                adCost=_fmt_fee(m.min_advertisement_fee_krw),
                saleType=_SALE_TYPE.get(m.sales_type, m.sales_type or "-"),
                updatedAt=_fmt_date(m.source_updated_at),
                createdAt=_fmt_date(m.source_created_at),
            )
        )
    return items


# ===== admin 매체 상세/등록/수정 (media 전 컬럼) =====


def _media_columns() -> list[str]:
    """Media 모델의 전체 컬럼명(선언 순서)."""
    return [c.key for c in sa_inspect(Media).mapper.column_attrs]


def _xlsx_cell(val):
    """openpyxl 셀에 넣을 수 있는 값으로 변환."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return "Y" if val else "N"
    if isinstance(val, Decimal):
        return float(val)
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    return val


def _build_xlsx(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "media"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_media_xlsx(db: Session) -> bytes:
    """전체 매체 데이터를 xlsx 로 export — 헤더=전 컬럼, 행=매체."""
    cols = _media_columns()
    medias = db.query(Media).order_by(Media.created_at.desc()).all()
    rows = [[_xlsx_cell(getattr(m, c)) for c in cols] for m in medias]
    return _build_xlsx(cols, rows)


def media_template_xlsx() -> bytes:
    """엑셀 일괄등록용 빈 양식 — 헤더=등록 폼 컬럼(자동 컬럼 제외)."""
    cols = [c for c in _media_columns() if c not in _MEDIA_AUTO_COLS]
    return _build_xlsx(cols, [])


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _coerce_import_value(col, raw):
    """xlsx 셀 값을 컬럼 타입에 맞는 파이썬 값으로 변환 — export(_xlsx_cell)의 역변환.

    셀은 openpyxl 네이티브 타입(int/float/bool/datetime) 또는 문자열로 들어온다.
    빈 값은 None. 변환 불가 시 ValueError 를 던져 호출부가 행 단위로 처리한다.
    """
    if _is_blank(raw):
        return None
    if isinstance(raw, str):
        raw = raw.strip()

    col_type = col.type
    if isinstance(col_type, Boolean):
        if isinstance(raw, bool):
            return raw
        s = str(raw).strip().lower()
        if s in ("y", "true", "1"):
            return True
        if s in ("n", "false", "0"):
            return False
        raise ValueError(f"Y/N 값이 올바르지 않습니다: {raw!r}")
    if isinstance(col_type, JSONB):
        if isinstance(raw, (dict, list)):
            return raw
        return json.loads(raw)
    if isinstance(col_type, DateTime):
        if isinstance(raw, datetime):
            return raw
        return datetime.fromisoformat(str(raw))
    if isinstance(col_type, (Integer, BigInteger)):
        if isinstance(raw, bool):
            raise ValueError(f"정수 값이 올바르지 않습니다: {raw!r}")
        if isinstance(raw, (int, float)):
            return int(raw)
        return int(str(raw))
    if isinstance(col_type, Numeric):
        return float(raw)
    return str(raw)


_IMPORT_MAX_ERRORS = 50


def import_media_xlsx(db: Session, content: bytes) -> dict:
    """엑셀 일괄등록 — media_id(=No) 기준 중복 제외, 없는 행만 삽입.

    반환: {total, inserted, skipped, failed, errors}. best-effort — 행별 savepoint 로
    한 행이 실패해도 나머지는 계속 삽입한다.
    """
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="엑셀 파일을 읽을 수 없습니다.")

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise HTTPException(status_code=400, detail="빈 파일입니다.")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    editable = set(_media_columns()) - _MEDIA_AUTO_COLS
    if "media_id" not in headers:
        raise HTTPException(
            status_code=400,
            detail="media_id 컬럼이 없습니다. 엑셀 양식(No 컬럼)을 확인해 주세요.",
        )

    col_index = {h: i for i, h in enumerate(headers) if h in editable}
    mid_idx = headers.index("media_id")
    columns = sa_inspect(Media).columns
    existing_ids = {r[0] for r in db.query(Media.media_id).all()}

    seen: set[str] = set()
    inserted = skipped = failed = 0
    errors: list[dict] = []

    def cell(row: tuple, idx: int):
        return row[idx] if idx < len(row) else None

    for row_no, row in enumerate(rows, start=2):
        if all(_is_blank(c) for c in row):
            continue

        mid_raw = cell(row, mid_idx)
        if _is_blank(mid_raw):
            failed += 1
            if len(errors) < _IMPORT_MAX_ERRORS:
                errors.append({"row": row_no, "media_id": "", "reason": "media_id(No) 누락"})
            continue
        mid = str(mid_raw).strip()

        if mid in existing_ids or mid in seen:
            skipped += 1
            continue

        data: dict = {}
        row_error: str | None = None
        for col_name, idx in col_index.items():
            try:
                data[col_name] = _coerce_import_value(columns[col_name], cell(row, idx))
            except (ValueError, TypeError, json.JSONDecodeError) as exc:
                row_error = f"{col_name}: {exc}"
                break
        if row_error:
            failed += 1
            if len(errors) < _IMPORT_MAX_ERRORS:
                errors.append({"row": row_no, "media_id": mid, "reason": row_error})
            continue

        if _is_blank(data.get("name")):
            failed += 1
            if len(errors) < _IMPORT_MAX_ERRORS:
                errors.append({"row": row_no, "media_id": mid, "reason": "매체명(name) 누락"})
            continue

        data["media_id"] = mid
        try:
            with db.begin_nested():
                db.add(Media(**data))
                db.flush()
            inserted += 1
            seen.add(mid)
        except Exception:
            failed += 1
            if len(errors) < _IMPORT_MAX_ERRORS:
                errors.append({"row": row_no, "media_id": mid, "reason": "저장 실패"})

    db.commit()
    return {
        "total": inserted + skipped + failed,
        "inserted": inserted,
        "skipped": skipped,
        "failed": failed,
        "errors": errors,
    }


def _image_dict(img: MediaImage) -> dict:
    return dict(
        id=str(img.id),
        image_url=img.image_url,
        sort_order=img.sort_order,
        is_thumbnail=img.is_thumbnail,
    )


def _serialize_media(m: Media) -> dict:
    """Media 전 컬럼 + 이미지 목록을 JSON 친화 dict 로 직렬화."""
    data: dict = {}
    for col in _media_columns():
        val = getattr(m, col)
        data[col] = float(val) if isinstance(val, Decimal) else val
    data["images"] = [_image_dict(img) for img in m.images]
    return data


def _get_media_or_404(db: Session, media_id: str) -> Media:
    m = db.query(Media).filter(Media.media_id == media_id).first()
    if m is None:
        raise HTTPException(status_code=404, detail="매체를 찾을 수 없습니다.")
    return m


def get_admin_media(db: Session, media_id: str) -> dict:
    return _serialize_media(_get_media_or_404(db, media_id))


def create_media(db: Session, payload: dict) -> dict:
    media_id = (payload.get("media_id") or "").strip()
    if not media_id:
        raise HTTPException(status_code=400, detail="매체 ID(media_id)는 필수입니다.")
    if db.query(Media).filter(Media.media_id == media_id).first():
        raise HTTPException(status_code=409, detail="이미 존재하는 매체 ID입니다.")
    editable = set(_media_columns()) - _MEDIA_AUTO_COLS
    data = {k: v for k, v in payload.items() if k in editable}
    data["media_id"] = media_id
    m = Media(**data)
    db.add(m)
    db.commit()
    db.refresh(m)
    return _serialize_media(m)


def update_media(db: Session, media_id: str, payload: dict) -> dict:
    m = _get_media_or_404(db, media_id)
    editable = set(_media_columns()) - _MEDIA_AUTO_COLS - {"media_id"}
    for key, value in payload.items():
        if key in editable:
            setattr(m, key, value)
    db.commit()
    db.refresh(m)
    return _serialize_media(m)


def delete_media(db: Session, media_id: str) -> None:
    db.delete(_get_media_or_404(db, media_id))
    db.commit()


def add_media_image(db: Session, media_id: str, file: UploadFile) -> dict:
    """업로드 파일을 저장하고 media_image 행 추가. 첫 이미지는 대표(is_thumbnail)로.

    thumbnail_url 컬럼은 건드리지 않는다(별도 텍스트 필드로 독립 관리).
    """
    m = _get_media_or_404(db, media_id)
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in _MEDIA_IMAGE_EXTS:
        raise HTTPException(status_code=400, detail="지원하지 않는 이미지 형식입니다.")
    content = file.file.read(_MEDIA_IMAGE_MAX_BYTES + 1)
    if len(content) > _MEDIA_IMAGE_MAX_BYTES:
        raise HTTPException(status_code=400, detail="이미지 용량은 10MB 이하만 가능합니다.")

    upload_dir = get_settings().upload_dir
    rel_dir = os.path.join("media", media_id)
    abs_dir = Path(upload_dir) / rel_dir
    abs_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}{ext}"
    (abs_dir / stored_name).write_bytes(content)
    image_url = f"/uploads/{rel_dir}/{stored_name}"

    next_order = max((img.sort_order for img in m.images), default=-1) + 1
    is_first = len(m.images) == 0
    db.add(
        MediaImage(
            media_id=media_id,
            image_url=image_url,
            sort_order=next_order,
            is_thumbnail=is_first,
        )
    )
    db.commit()
    db.refresh(m)
    return _serialize_media(m)


def delete_media_image(db: Session, media_id: str, image_id: str) -> dict:
    m = _get_media_or_404(db, media_id)
    img = (
        db.query(MediaImage)
        .filter(MediaImage.id == image_id, MediaImage.media_id == media_id)
        .first()
    )
    if img is None:
        raise HTTPException(status_code=404, detail="이미지를 찾을 수 없습니다.")
    was_thumb = img.is_thumbnail
    db.delete(img)
    db.flush()
    if was_thumb:
        remaining = (
            db.query(MediaImage)
            .filter(MediaImage.media_id == media_id)
            .order_by(MediaImage.sort_order)
            .first()
        )
        if remaining:
            remaining.is_thumbnail = True
    db.commit()
    db.refresh(m)
    return _serialize_media(m)
