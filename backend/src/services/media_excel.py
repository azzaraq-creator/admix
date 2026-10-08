"""어드민 매체 엑셀 — 다운로드·업로드 양식·일괄 등록을 어드민 매체 폼과 같은 구조로.

- 칸 순서·이름·구역은 어드민 폼(frontend/app/admin/(main)/media/_components/mediaFields.ts)과 같다.
  폼을 고치면 여기도 함께 고친다(tests/test_media_excel.py 가 빠진 컬럼을 잡는다).
- 폼의 "사용하지 않는 값" 구역(UNUSED_COLUMNS)은 내보내지도 받지도 않는다.
- 선택값은 한글로(INSIDE → 공간형), 예/아니오, 태그는 쉼표로 풀어 쓴다. 올릴 때는 거꾸로 바꾼다.
- 상품(media_plan)은 "상품" 시트에 한 줄씩. 다운로드는 매체 ID로, 업로드 양식은 "연결 번호"로 매체와 잇는다.
- 일괄 등록은 이 한글 양식(1행 구역, 2행 칸 이름)과 예전 영문 양식(1행 DB 컬럼명)을 모두 받는다.
"""
from __future__ import annotations

import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from sqlalchemy import inspect as sa_inspect
from sqlalchemy.orm import Session, selectinload

from src.models.media_master import Media
from src.services import kakao_address
from src.services import media_service as ms

# (컬럼, 칸 이름, 값 변환) — 변환: None(그대로) | dict(선택값→한글) | "bool" | "fee_yn" | "tags" | "won"
#   | "formats"(소재 형식 코드 목록 ↔ "영상(MP4), 스틸컷(JPG·PNG)") | "time"("06:00")
_YES_NO = "bool"

SECTIONS: list[tuple[str, list[tuple[str, str, object]]]] = [
    (
        "기본 정보",
        [
            ("media_id", "매체 ID", None),
            ("name", "매체명", None),
            ("second_name", "구분명", None),
            ("media_source", "고정/이동", {"FIXED": "고정", "MOVING": "이동"}),
            ("sales_type", "판매 유형", {"SINGLE": "단품", "GROUP": "묶음"}),
            ("category_large", "카테고리(대)", None),
            ("category_small", "카테고리(소)", None),
            ("ooh_type", "매체 타입", None),
            ("exposure_type", "설치 장소", {"INSIDE": "공간형", "OUTSIDE": "외부형"}),
            (
                "media_shape",
                "매체 형태",
                {
                    "HORIZONTAL_SHAPE": "가로형",
                    "VERTICAL_SHAPE": "세로형",
                    "CURVED_SHAPE": "커브형",
                    "DIFFERENT_SHAPE": "변형",
                },
            ),
            ("description", "매체 설명", None),
        ],
    ),
    (
        "위치 · 운행 지역",
        [
            ("address", "도로명 주소", None),
            ("accurate_address", "도로명 주소(동 포함)", None),
            ("full_address_jibun", "지번 주소", None),
            ("building_name", "건물명", None),
            ("legal_dong", "법정동", None),
            ("city", "시/도 (이동: 운행 시·도)", None),
            ("district", "구/군 (이동: 운행 구)", None),
            ("address_detail", "상세 주소", None),
            ("latitude", "위도", None),
            ("longitude", "경도", None),
            ("moving_location_detail", "노선·운행 설명 (이동)", None),
        ],
    ),
    (
        "가격·판매",
        [
            ("min_advertisement_fee_krw", "최소 광고비(원)", "won"),
            ("max_advertisement_fee_krw", "최대 광고비(원)", "won"),
            ("any_production_fee_yn", "제작비 유무", "fee_yn"),
            ("min_production_fee_krw", "최소 제작비(원)", "won"),
            ("max_production_fee_krw", "최대 제작비(원)", "won"),
            ("execution_status", "집행 상태", None),
            ("lead_time_bizdays", "리드타임(영업일)", None),
            ("plan_count", "상품(플랜) 수", None),
        ],
    ),
    (
        "규격·수량",
        [
            ("device_quantity", "기기 수량", None),
            ("surface_quantity", "면 수량", None),
            ("spec_width", "규격 가로(m)", None),
            ("spec_height", "규격 세로(m)", None),
            ("spec_unit", "규격 단위", None),  # 다운로드 참고용 — 항상 m, 양식에는 없다
            ("spec_resolution_width", "해상도 가로(px)", None),
            ("spec_resolution_height", "해상도 세로(px)", None),
        ],
    ),
    (
        "소재·운영 시간",
        [
            ("material_formats", "소재 형식", "formats"),
            ("operation_start_time", "운영 시작", "time"),
            ("operation_end_time", "운영 종료", "time"),
        ],
    ),
    (
        "월평균 유동인구(직접 입력)",
        [
            ("population_count", "월평균 유동인구(명)", None),
            ("population_note", "유동인구 기준", None),
            ("population_male_pct", "남성 비율(%)", None),
            ("population_female_pct", "여성 비율(%)", None),
            ("population_age_10", "10대 이하(%)", None),
            ("population_age_20", "20대(%)", None),
            ("population_age_30", "30대(%)", None),
            ("population_age_40", "40대(%)", None),
            ("population_age_50", "50대(%)", None),
            ("population_age_60", "60대 이상(%)", None),
        ],
    ),
    (
        "노출·검색",
        [
            ("is_popular_yn", "인기 매체", _YES_NO),
            ("is_newly_built_yn", "신규 매체", _YES_NO),
            ("popular_type", "인기 업종", "tags"),
            ("loc_label", "지역 라벨", None),
            ("list_labels", "검색 태그", "tags"),
            ("special_remarks", "판매 메모(특이사항)", None),
        ],
    ),
    (
        "등급·추천",
        [
            ("final_grade", "최종 등급", None),
            ("quality_score", "품질 점수", None),
            ("gangnam_dong_grade", "강남 동별 등급", None),
            ("grade_method", "등급 산정 방식", None),
            ("audience_summary", "타깃(유동인구) 요약", None),
        ],
    ),
    (
        "상권",
        [
            ("market_area", "상권(호칭)", None),
            ("loc_code", "지역 코드", None),
            ("market_dong", "상권 세부(동)", None),
            ("market_keyword", "원천 상권명", None),
            ("market_profile_id", "상권 프로파일 ID", None),
        ],
    ),
    (
        "원천 데이터",
        [
            ("source_detail_id", "원본 상세 ID", None),
            ("category_id", "카테고리 ID", None),
            ("parent_category_code", "상위 카테고리 코드", None),
            ("children_count", "하위 매체 수", None),
            ("company_media_id", "업체 매체 ID", None),
            ("company_mapper_user_id", "업체 매퍼 사용자 ID", None),
            ("source_created_at", "원본 생성일시", None),
            ("source_updated_at", "원본 수정일시", None),
            ("road_view_latitude", "로드뷰 위도", None),
            ("road_view_longitude", "로드뷰 경도", None),
            ("road_view_heading", "로드뷰 방향", None),
            ("road_view_pitch", "로드뷰 피치", None),
            ("markers_vo", "마커 정보(JSON)", None),
            ("map_bounds_vo", "지도 경계(JSON)", None),
            ("recommended_media_items", "추천 매체(JSON)", None),
            ("created_at", "등록일시", None),
            ("updated_at", "수정일시", None),
        ],
    ),
]

# 어드민 폼 "사용하지 않는 값" 구역 — 사이트·추천에 쓰이지 않아 내보내지 않는다.
UNUSED_COLUMNS = frozenset(
    {
        "device_type",
        "feature",
        "production_fees_summary",
        "media_shape_summary",
        "properties_count",
        "properties_type",
        "properties_summary",
        "properties_extra_json",
        "thumbnail_url",
        "image_count",
        "gangnam_grade_reason",
        "grade_evidence",
        "area_evidence",
        "ind_evidence",
    }
)

# 내보낼 때 덧붙이는 계산 칸(사진에서 만든다).
_PHOTO_SECTION = ("사진", ["대표 이미지 URL", "사진 수"])

# "상품" 시트 칸 — 어드민 매체 폼의 상품 칸과 같다(media_service._PLAN_EDITABLE).
PLAN_MASTER_TYPES = {"PM_INDIVIDUAL": "개별", "PM_PACKAGE": "패키지", "PM_NETWORK": "네트워크"}
PLAN_DURATION_TYPES = {"MONTHS": "개월", "WEEKS": "주", "YEARS": "년", "DAYS": "일"}
PLAN_COLUMNS: list[tuple[str, str, object]] = [
    ("product_display_name", "상품명", None),
    ("product_master_type", "판매 방식", PLAN_MASTER_TYPES),
    ("contractual_duration", "계약 기간", None),
    ("contractual_duration_type", "기간 단위", PLAN_DURATION_TYPES),
    ("advertisement_fee", "광고비(원)", "won"),
    ("production_fee", "제작비(원)", "won"),
    ("exposure_duration_seconds", "노출 시간(초)", None),
    ("broadcasts_count_manual", "일 송출 수", None),
    ("default_device_quantity", "기기 수량", None),
    ("default_surface_quantity", "면 수량", None),
]
PLAN_REQUIRED = frozenset(
    {"product_display_name", "product_master_type", "contractual_duration", "contractual_duration_type"}
)
LINK_LABEL = "연결 번호"
PLAN_SHEET = "상품"


def _style_header_row(ws, row: int, labels: list[str], required: frozenset | set = frozenset()) -> None:
    """한 줄 머리글 — 필수 칸은 * 와 빨간색."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for col, label in enumerate(labels, start=1):
        need = label in required
        cell = ws.cell(row=row, column=col, value=f"{label} *" if need else label)
        cell.font = Font(bold=True, color="D65856" if need else None)
        cell.fill = PatternFill("solid", fgColor="FDECEC" if need else "F5F5F5")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = max(12, len(label) * 2 + 4)


def _write_plan_sheet(wb, medias: list[Media]) -> None:
    """다운로드의 "상품" 시트 — 한 줄에 상품 하나, 매체 ID·상품 번호와 함께(상품 번호 순)."""
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(PLAN_SHEET)
    fixed = ["매체 ID", "매체명", "상품 번호"]
    _style_header_row(ws, 1, [*fixed, *[c[1] for c in PLAN_COLUMNS]])
    won_cols = {
        len(fixed) + i for i, (_, _, conv) in enumerate(PLAN_COLUMNS, start=1) if conv == "won"
    }
    row = 2
    for m in medias:
        name = " ".join(p for p in [(m.name or "").strip(), (m.second_name or "").strip()] if p)
        for plan in sorted(m.plans, key=lambda p: p.plan_no):
            values = [m.media_id, name, plan.plan_no]
            for key, _, conv in PLAN_COLUMNS:
                v = getattr(plan, key)
                if key == "contractual_duration_type" and v and not v.endswith("S"):
                    v = f"{v}S"  # 원천에는 단수형(MONTH)도 있다
                values.append(_cell(v, conv))
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                if col in won_cols:
                    cell.number_format = "#,##0"
            row += 1
    ws.freeze_panes = "D2"
    last_col = get_column_letter(len(fixed) + len(PLAN_COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, row - 1)}"


def _cell(value, conv):
    if value is None:
        return ""
    if conv == _YES_NO:
        return "예" if value else "아니오"
    if conv == "fee_yn":
        return "있음" if value else "없음"
    if conv == "tags":
        items = value if isinstance(value, list) else str(value).replace("|", ",").split(",")
        return ", ".join(str(v).strip() for v in items if str(v).strip())
    if conv == "formats":
        return ", ".join(ms.MATERIAL_FORMATS.get(v, v) for v in value)
    if isinstance(conv, dict):
        return conv.get(value, value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _thumbnail(m: Media) -> str:
    """대표 이미지 — 대표 표시 우선, 없으면 첫 사진(기획안과 같은 기준)."""
    imgs = sorted(m.images, key=lambda i: (not i.is_thumbnail, i.sort_order))
    return imgs[0].image_url if imgs else ""


def export_media_xlsx(db: Session) -> bytes:
    """전체 매체 — "매체" 시트(1행 구역 이름, 2행 칸 이름, 3행부터 매체·매체 ID 순)와 "상품" 시트."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    columns = [col for _, cols in SECTIONS for col in cols]
    medias = (
        db.query(Media)
        .options(selectinload(Media.images), selectinload(Media.plans))
        .order_by(Media.media_id)
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "매체"

    section_fill = PatternFill("solid", fgColor="EDE3F7")
    header_fill = PatternFill("solid", fgColor="F5F5F5")
    bold = Font(bold=True)
    thin = Side(style="thin", color="D9D9D9")
    center = Alignment(horizontal="center", vertical="center")

    # 1행: 구역 이름(병합), 2행: 칸 이름
    col_no = 1
    for title, cols in [*SECTIONS, (_PHOTO_SECTION[0], _PHOTO_SECTION[1])]:
        labels = [c[1] for c in cols] if title != _PHOTO_SECTION[0] else cols
        start = col_no
        for label in labels:
            cell = ws.cell(row=2, column=col_no, value=label)
            cell.font, cell.fill, cell.alignment = bold, header_fill, center
            cell.border = Border(bottom=thin)
            col_no += 1
        ws.cell(row=1, column=start, value=title)
        if col_no - 1 > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col_no - 1)
        head = ws.cell(row=1, column=start)
        head.font, head.fill, head.alignment = bold, section_fill, center

    won_cols = [i for i, (_, _, conv) in enumerate(columns, start=1) if conv == "won"]
    for row_no, m in enumerate(medias, start=3):
        for i, (key, _, conv) in enumerate(columns, start=1):
            ws.cell(row=row_no, column=i, value=_cell(getattr(m, key), conv))
        ws.cell(row=row_no, column=len(columns) + 1, value=_thumbnail(m))
        ws.cell(row=row_no, column=len(columns) + 2, value=len(m.images))
        for i in won_cols:
            ws.cell(row=row_no, column=i).number_format = "#,##0"

    # 칸 너비 — 칸 이름 길이 기준(긴 글 칸은 넓게), 첫 두 칸(ID·매체명)과 머리 두 줄 고정.
    wide = {"description", "audience_summary", "special_remarks", "markers_vo",
            "map_bounds_vo", "recommended_media_items", "list_labels"}
    for i, (key, label, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 40 if key in wide else max(12, len(label) * 2 + 2)
    ws.column_dimensions[get_column_letter(len(columns) + 1)].width = 40
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns) + 2)}{max(2, len(medias) + 2)}"

    _write_plan_sheet(wb, medias)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 업로드 양식 · 일괄 등록 ─────────────────────────────────────────

# 양식에 넣지 않는 칸 — 매체 ID(자동 부여), 다른 데이터에서 정해지는 값.
# 광고비·제작비 범위와 제작비 유무는 "상품" 시트에서 계산한다.
# 상권·원천 데이터 구역(원천 시스템이 계산해 둔 값)도 양식에서는 뺀다. "사용하지 않는 값"은 SECTIONS 에 없다.
# 다운로드 엑셀을 그대로 올리면 이 칸들도 읽는다(상품에서 계산한 값이 덮어쓴다).
_PLAN_DERIVED = frozenset(
    {
        "min_advertisement_fee_krw",
        "max_advertisement_fee_krw",
        "min_production_fee_krw",
        "max_production_fee_krw",
        "any_production_fee_yn",
        "plan_count",
    }
)
# 고정 매체 위치 칸 — 도로명 주소로 카카오 주소 검색을 해서 비어 있는 칸을 채운다(적어 둔 값은 그대로 둔다).
# 양식에서는 이 칸들을 빼고 도로명 주소·상세 주소만 받는다. 시/도·구/군은 이동 매체 운행 지역 때문에 남긴다.
_ADDRESS_FILLED = (
    "accurate_address",
    "full_address_jibun",
    "building_name",
    "legal_dong",
    "city",
    "district",
    "latitude",
    "longitude",
)
_TEMPLATE_ADDRESS_SKIP = frozenset(_ADDRESS_FILLED) - {"city", "district"}
# 양식에서만 쓰는 칸 이름 — 고정 매체는 자동으로 채우니 이동 매체 칸으로 보이게 한다(다운로드 엑셀은 원래 이름).
_TEMPLATE_LABELS = {"city": "시/도 (이동 매체)", "district": "구/군 (이동 매체)"}
_TEMPLATE_SKIP = (
    frozenset({"media_id", "created_at", "updated_at", "spec_unit"}) | _PLAN_DERIVED | _TEMPLATE_ADDRESS_SKIP
)
_TEMPLATE_SKIP_SECTIONS = frozenset({"상권", "원천 데이터"})

# 필수 칸 — 어드민 폼과 같다(목록 카드·필터·지도에 꼭 필요한 값). 가격은 상품에서 온다.
REQUIRED = frozenset(
    {
        "name",
        "media_source",
        "sales_type",
        "category_large",
        "category_small",
        "ooh_type",
        "exposure_type",
    }
)
REQUIRED_FIXED = frozenset({"address", "latitude", "longitude"})
REQUIRED_MOVING = frozenset({"city"})

# 드롭다운 선택지(값이 정해진 칸) — 위 SECTIONS 의 한글 선택값과 같다.
_GRADES = ["S", "A", "B", "C", "D"]
_CHOICES: dict[str, list[str]] = {
    "ooh_type": ["OOH", "DOOH"],
    "execution_status": ["가능", "확인 필요", "불가"],
    "final_grade": _GRADES,
    "gangnam_dong_grade": _GRADES,
}

# 작성 안내 — (칸, 입력 방법, 예시)
_GUIDE: dict[str, tuple[str, str]] = {
    "name": ("장소·매체 이름. 화면에는 '매체명 구분명'으로 붙어 보입니다.", "신사역"),
    "second_name": ("같은 매체명이 여러 개일 때 구분하는 말. 없으면 비웁니다.", "벽면"),
    "media_source": ("고정 / 이동 중에서 고릅니다.", "고정"),
    "sales_type": ("단품 / 묶음", "단품"),
    "category_large": ("기존 카테고리를 쓰면 매체 찾기 필터가 맞게 묶입니다.", "전광판/빌보드"),
    "category_small": ("카테고리(대)에 속한 소분류.", "전광판"),
    "ooh_type": ("OOH(지면·실물) / DOOH(디지털). DOOH는 상품 제작비가 없는 것으로 저장됩니다.", "DOOH"),
    "exposure_type": ("공간형(실내 공간 안) / 외부형(건물 밖·거리)", "외부형"),
    "media_shape": ("가로형 / 세로형 / 커브형 / 변형", "가로형"),
    "address": (
        "고정 매체 필수. 도로명 주소만 적으면 올릴 때 도로명 주소(동 포함)·지번 주소·건물명·법정동·"
        "시/도·구/군·위도·경도를 카카오 주소 검색으로 자동 입력합니다.",
        "서울 강남구 강남대로 396",
    ),
    "address_detail": ("층·출구 등 자세한 위치. 없으면 비웁니다.", "2층 로비"),
    "city": ("이동 매체 필수. 운행 시·도(서울특별시, 경기도 … 또는 전국). 고정 매체는 비웁니다(자동 입력).", "서울특별시"),
    "district": ("이동 매체의 운행 구를 쉼표로(비우면 시·도 전역). 고정 매체는 비웁니다(자동 입력).", "강남구, 서초구"),
    "moving_location_detail": ("이동 매체의 노선·운행 설명.", "146번 간선버스 · 상계동 ~ 강남역"),
    "execution_status": ("가능 / 확인 필요 / 불가", "가능"),
    "lead_time_bizdays": ("신청부터 송출까지 영업일.", "7"),
    "device_quantity": ("기획안의 수량(기). 음수 불가.", "1"),
    "surface_quantity": ("기획안의 면. 음수 불가.", "1"),
    "spec_width": ("가로 길이(m, 숫자).", "10"),
    "spec_height": ("세로 길이(m, 숫자).", "5"),
    "spec_resolution_width": ("화면 해상도 가로(px). 매체 타입이 DOOH일 때만 저장합니다.", "1920"),
    "spec_resolution_height": ("화면 해상도 세로(px). 매체 타입이 DOOH일 때만 저장합니다.", "1080"),
    "material_formats": (
        "쉼표로 여러 개: " + " / ".join(ms.MATERIAL_FORMATS.values()),
        "영상(MP4), 스틸컷(JPG·PNG)",
    ),
    "operation_start_time": ("DOOH 운영 시작 시각(HH:MM). 매체 타입이 DOOH일 때만 저장합니다.", "06:00"),
    "operation_end_time": ("DOOH 운영 종료 시각(HH:MM, 자정은 24:00).", "24:00"),
    "population_count": ("월평균 유동인구. 실시간 인구(서울시 주요 121장소)를 가져올 수 없는 고정 매체에 적으면 원천 상권 데이터보다 먼저 보입니다. 이동 매체는 쓰지 않습니다.", "450000"),
    "population_note": ("유동인구 수치의 기준(시기·출처). 팝업에 함께 보입니다.", "2025년 3분기"),
    "population_male_pct": ("남성 비율(%). 여성과 합쳐 100.", "48"),
    "population_female_pct": ("여성 비율(%).", "52"),
    "population_age_10": ("연령대 비율(%) — 여섯 칸 합쳐 100.", "8"),
    "population_age_60": ("60대 이상 비율(%).", "12"),
    "is_popular_yn": ("예 / 아니오 — 카드에 '인기' 뱃지", "아니오"),
    "is_newly_built_yn": ("예 / 아니오 — 카드에 '신규' 뱃지", "예"),
    "popular_type": ("쉼표로 여러 개.", "패션, 화장품"),
    "loc_label": ("이 말로 검색해도 매체가 나옵니다.", "강남역"),
    "list_labels": ("검색 태그. 쉼표로 여러 개.", "강남, 대형, 전광판"),
    "quality_score": ("0~100", "80"),
    "audience_summary": ("믹시가 읽는 형식을 지켜 주세요.", "일평균 45.8만 · 여성 47% · 20·40대 중심"),
}
_PLAN_GUIDE: dict[str, tuple[str, str]] = {
    "product_display_name": ("상품 이름. 매체 상세·기획안에 보입니다.", "영상(20초)"),
    "product_master_type": ("개별 / 패키지 / 네트워크 — 매체 찾기 '판매 유형' 필터에 쓰입니다.", "개별"),
    "contractual_duration": ("1 이상 정수.", "1"),
    "contractual_duration_type": ("개월 / 주 / 년 / 일", "개월"),
    "advertisement_fee": ("계약 기간 동안의 광고비(원). 비우면 가격 미정(-).", "4000000"),
    "production_fee": ("1회 제작비(원). 없으면 비웁니다. 매체 타입이 DOOH면 저장하지 않습니다.", "200000"),
    "exposure_duration_seconds": ("영상 상품의 한 번 송출 길이. 매체 타입이 OOH면 저장하지 않습니다.", "20"),
    "broadcasts_count_manual": ("하루 송출 횟수. 매체 타입이 OOH면 저장하지 않습니다.", "100"),
    "default_device_quantity": ("기기 수량. 비우면 1.", "1"),
    "default_surface_quantity": ("면 수량. 비우면 1.", "1"),
}


def _template_sections() -> list[tuple[str, list[tuple[str, str, object]]]]:
    return [
        (title, [(k, _TEMPLATE_LABELS.get(k, label), conv) for k, label, conv in cols if k not in _TEMPLATE_SKIP])
        for title, cols in SECTIONS
        if title not in _TEMPLATE_SKIP_SECTIONS
    ]


def _is_required(key: str) -> str | None:
    """필수 표시 — 항상이면 "필수", 고정·이동 한쪽만이면 그 이름."""
    if key in REQUIRED:
        return "필수"
    if key in REQUIRED_FIXED:
        return "고정 매체 필수"
    if key in REQUIRED_MOVING:
        return "이동 매체 필수"
    return None


def _choices(key: str, conv) -> list[str] | None:
    if isinstance(conv, dict):
        return list(conv.values())
    if conv == _YES_NO:
        return ["예", "아니오"]
    if conv == "fee_yn":
        return ["있음", "없음"]
    return _CHOICES.get(key)


def _add_dropdown(ws, col: int, options: list[str], first_row: int) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="선택지에서 고르세요",
        error=" / ".join(options),
    )
    ws.add_data_validation(dv)
    letter = get_column_letter(col)
    dv.add(f"{letter}{first_row}:{letter}{first_row + 999}")


def media_template_xlsx() -> bytes:
    """일괄 등록 양식 — "매체 등록"·"상품"·"작성 안내" 시트.

    매체 등록: 다운로드 엑셀과 같은 구조(1행 구역, 2행 칸 이름, 필수는 *), 맨 앞에 연결 번호.
    상품: 한 줄에 상품 하나, 연결 번호로 매체와 잇는다. 선택 칸은 드롭다운.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sections = _template_sections()
    wb = Workbook()
    ws = wb.active
    ws.title = "매체 등록"
    section_fill = PatternFill("solid", fgColor="EDE3F7")
    header_fill = PatternFill("solid", fgColor="F5F5F5")
    required_fill = PatternFill("solid", fgColor="FDECEC")
    bold, red = Font(bold=True), Font(bold=True, color="D65856")
    center = Alignment(horizontal="center", vertical="center")

    # 1열: 연결 번호(상품 시트와 잇는 값)
    link = ws.cell(row=2, column=1, value=f"{LINK_LABEL} *")
    link.font, link.fill, link.alignment = red, required_fill, center
    head = ws.cell(row=1, column=1, value="연결")
    head.font, head.fill, head.alignment = bold, section_fill, center
    ws.column_dimensions["A"].width = 12

    col_no = 2
    for title, cols in sections:
        start = col_no
        for key, label, conv in cols:
            need = _is_required(key)
            cell = ws.cell(row=2, column=col_no, value=f"{label} *" if need else label)
            cell.font = red if need else bold
            cell.fill = required_fill if need else header_fill
            cell.alignment = center
            ws.column_dimensions[get_column_letter(col_no)].width = max(14, len(label) * 2 + 4)
            options = _choices(key, conv)
            if options:
                _add_dropdown(ws, col_no, options, 3)
            col_no += 1
        ws.cell(row=1, column=start, value=title)
        if col_no - 1 > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col_no - 1)
        head = ws.cell(row=1, column=start)
        head.font, head.fill, head.alignment = bold, section_fill, center
    ws.freeze_panes = "C3"

    plans_ws = wb.create_sheet(PLAN_SHEET)
    plan_labels = [LINK_LABEL, *[c[1] for c in PLAN_COLUMNS]]
    required_labels = {LINK_LABEL} | {label for key, label, _ in PLAN_COLUMNS if key in PLAN_REQUIRED}
    _style_header_row(plans_ws, 1, plan_labels, required_labels)
    for col, (_, _, conv) in enumerate(PLAN_COLUMNS, start=2):
        if isinstance(conv, dict):
            _add_dropdown(plans_ws, col, list(conv.values()), 2)
    plans_ws.freeze_panes = "B2"

    guide = wb.create_sheet("작성 안내")
    guide.append(["시트", "칸 이름", "필수", "입력 방법", "예시"])
    for cell in guide[1]:
        cell.font, cell.fill = bold, header_fill
    guide.append([
        "(공통)", "", "",
        "매체 등록 시트 3행부터 한 줄에 매체 하나, 상품 시트 2행부터 한 줄에 상품 하나를 적습니다. "
        "두 시트의 '연결 번호'가 같으면 그 매체의 상품입니다(매체마다 상품 하나 이상 필수). "
        "매체 ID는 저장할 때 자동으로 매겨지고(고정 F000001 · 이동 M000001), "
        "광고비·제작비 범위는 상품에서 계산됩니다. 사진은 등록 후 어드민 매체 화면에서 올립니다.",
        "",
    ])
    guide.append(["매체 등록", LINK_LABEL, "필수", "상품 시트와 잇는 번호. 매체마다 다르게(1, 2, 3 …).", "1"])
    for _, cols in sections:
        for key, label, conv in cols:
            how, example = _GUIDE.get(key, ("", ""))
            options = _choices(key, conv)
            if options and not how:
                how = " / ".join(options)
            guide.append(["매체 등록", label, _is_required(key) or "", how, example])
    guide.append([PLAN_SHEET, LINK_LABEL, "필수", "이 상품이 속한 매체의 연결 번호.", "1"])
    for key, label, _ in PLAN_COLUMNS:
        how, example = _PLAN_GUIDE.get(key, ("", ""))
        guide.append([PLAN_SHEET, label, "필수" if key in PLAN_REQUIRED else "", how, example])
    for letter, width in zip("ABCDE", (12, 24, 14, 80, 36)):
        guide.column_dimensions[letter].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 바뀌기 전 칸 이름·선택값 — 예전에 받아 둔 엑셀을 올려도 읽히게 함께 받는다.
_OLD_LABELS = {"OOH 유형": "매체 타입"}
_OLD_VALUES = {
    "exposure_type": {"실내": "INSIDE", "실외": "OUTSIDE"},
    "media_shape": {"곡면형": "CURVED_SHAPE", "비정형": "DIFFERENT_SHAPE"},
}


def _header_key_map() -> dict[str, tuple[str, object]]:
    """칸 이름(한글) → (컬럼, 변환). 다운로드 엑셀을 그대로 올려도 읽히게 전 구역을 받는다."""
    labels = {label: (key, conv) for _, cols in SECTIONS for key, label, conv in cols}
    for old, new in _OLD_LABELS.items():
        labels[old] = labels[new]
    for key, label in _TEMPLATE_LABELS.items():
        labels[label] = labels[_label_of(key)]
    return labels


def _label_of(key: str) -> str:
    return next(label for _, cols in SECTIONS for k, label, _ in cols if k == key)


def _clean_label(v) -> str:
    return str(v or "").strip().rstrip("*").strip()


def is_korean_layout(rows: list[tuple]) -> bool:
    """1행 구역 이름·2행 한글 칸 이름 양식인지 — 2행에 아는 칸 이름이 있으면 그렇다."""
    if len(rows) < 2:
        return False
    labels = _header_key_map()
    return any(_clean_label(v) in labels for v in rows[1])


def _choice_value(conv: dict, text: str, old: dict | None = None) -> str:
    """한글 선택값(또는 코드) → 코드."""
    reverse = {v: k for k, v in conv.items()} | (old or {})
    if text in reverse:
        return reverse[text]
    if text in conv:  # 코드로 적어도 받는다(SINGLE 등)
        return text
    raise ValueError(f"{' / '.join(conv.values())} 중에서 골라 주세요: {text}")


def _parse_cell(key: str, conv, raw):
    """한글 양식 값 → DB 값. 바꿀 수 없으면 ValueError(행 오류로 보고한다)."""
    if ms._is_blank(raw):
        return None
    text = str(raw).strip()
    if isinstance(conv, dict):
        return _choice_value(conv, text, _OLD_VALUES.get(key))
    if conv in (_YES_NO, "fee_yn"):
        yes, no = ({"예", "Y", "y", "TRUE", "true", "1"}, {"아니오", "N", "n", "FALSE", "false", "0"})
        if conv == "fee_yn":
            yes, no = yes | {"있음"}, no | {"없음"}
        if text in yes:
            return True
        if text in no:
            return False
        raise ValueError(f"{'있음 / 없음' if conv == 'fee_yn' else '예 / 아니오'} 중에서 골라 주세요: {text}")
    if conv == "formats":
        reverse = {v: k for k, v in ms.MATERIAL_FORMATS.items()}
        codes = []
        for t in text.split(","):
            t = t.strip()
            if not t:
                continue
            code = reverse.get(t) or (t if t in ms.MATERIAL_FORMATS else None)
            if code is None:
                raise ValueError(f"{' / '.join(ms.MATERIAL_FORMATS.values())} 중에서 골라 주세요: {t}")
            if code not in codes:
                codes.append(code)
        return codes or None
    if conv == "time":
        # 엑셀이 "06:00"을 시각 값으로 바꿔 저장해도 읽는다("24:00"은 글자로 남는다).
        if hasattr(raw, "strftime"):
            return raw.strftime("%H:%M")
        hh, _, mm = text.partition(":")
        if not (hh.isdigit() and mm[:2].isdigit()):
            raise ValueError(f"HH:MM 형식으로 적어 주세요: {text}")
        return f"{int(hh):02d}:{mm[:2]}"
    if conv == "tags":
        items = []
        for t in text.replace("|", ",").split(","):
            t = t.strip()
            if t and t not in items:
                items.append(t)
        if not items:
            return None
        return items if key == "list_labels" else ",".join(items)
    if conv == "won" and isinstance(raw, str):
        raw = raw.replace(",", "").replace("원", "").strip()
    options = _CHOICES.get(key)
    if options and text not in options:
        raise ValueError(f"{' / '.join(options)} 중에서 골라 주세요: {text}")
    return ms._coerce_import_value(sa_inspect(Media).columns[key], raw)


# 상품 시트에서 글자 그대로 받는 칸(나머지는 선택값이거나 정수).
_PLAN_TEXT = frozenset({"product_display_name"})


def _parse_plan_cell(key: str, conv, raw):
    """상품 시트 값 → 서버(media_service._clean_plan)가 받는 값. 숫자는 정수로."""
    if ms._is_blank(raw):
        return None
    text = str(raw).strip()
    if key in _PLAN_TEXT:
        return text
    if isinstance(conv, dict):
        return _choice_value(conv, text)
    text = text.replace(",", "").replace("원", "").strip()
    try:
        number = float(text)
    except ValueError:
        raise ValueError(f"숫자를 입력해 주세요: {raw}") from None
    if number != int(number):
        raise ValueError(f"정수를 입력해 주세요: {raw}")
    return int(number)


def _link_key(raw) -> str:
    """연결 번호 — 1 과 1.0, "1" 을 같은 번호로 본다."""
    if ms._is_blank(raw):
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw).strip()


def _read_plans(wb, link_label: str) -> tuple[dict[str, list[tuple[int, dict]]], list[dict]]:
    """상품 시트 → {연결 값: [(행, 상품)]}, 오류 목록. 상품 시트가 없으면 빈 값."""
    if PLAN_SHEET not in wb.sheetnames:
        return {}, []
    rows = list(wb[PLAN_SHEET].iter_rows(values_only=True))
    if not rows:
        return {}, []
    labels = {label: (key, conv) for key, label, conv in PLAN_COLUMNS}
    header = [_clean_label(v) for v in rows[0]]
    if link_label not in header:
        return {}, [{"row": 1, "media_id": "", "reason": f"상품 시트: '{link_label}' 칸이 없습니다."}]
    link_idx = header.index(link_label)
    cols = [(i, *labels[h]) for i, h in enumerate(header) if h in labels]
    out: dict[str, list[tuple[int, dict]]] = {}
    errors: list[dict] = []
    for row_no, row in enumerate(rows[1:], start=2):
        if all(ms._is_blank(v) for v in row):
            continue
        link = _link_key(row[link_idx] if link_idx < len(row) else None)
        if not link:
            errors.append({"row": row_no, "media_id": "", "reason": f"상품 시트: {link_label}이 비었습니다."})
            continue
        plan: dict = {"plan_no": None}
        try:
            for i, key, conv in cols:
                plan[key] = _parse_plan_cell(key, conv, row[i] if i < len(row) else None)
        except ValueError as exc:
            errors.append({"row": row_no, "media_id": "", "reason": f"상품 시트: {exc}"})
            continue
        # 기기·면 수량은 비우면 1(어드민 폼의 기본값과 같다).
        for key in ("default_device_quantity", "default_surface_quantity"):
            if plan.get(key) is None:
                plan[key] = 1
        missing = [label for key, label, _ in PLAN_COLUMNS if key in PLAN_REQUIRED and plan.get(key) is None]
        if missing:
            errors.append({"row": row_no, "media_id": "", "reason": "상품 시트: 필수 칸이 비었습니다: " + ", ".join(missing)})
            continue
        out.setdefault(link, []).append((row_no, plan))
    return out, errors


def _fill_address(data: dict, cache: dict[str, dict | str]) -> str | None:
    """고정 매체의 비어 있는 위치 칸을 도로명 주소로 채운다. 실패하면 행 오류 문구를 돌려준다.

    적어 둔 값은 덮어쓰지 않는다. 위도·경도를 직접 적었으면 검색이 실패해도 그대로 넣는다.
    같은 주소는 한 번만 검색한다(cache: 주소 → 결과 dict 또는 오류 문구).
    """
    if data.get("media_source") == "MOVING":
        return None
    address = data.get("address")
    if ms._is_blank(address):
        return None
    # 건물명은 원래 없는 주소가 많아 검색할지 정하는 데 쓰지 않는다.
    if not any(ms._is_blank(data.get(k)) for k in _ADDRESS_FILLED if k != "building_name"):
        return None
    address = str(address).strip()
    if address not in cache:
        try:
            cache[address] = kakao_address.lookup(address)
        except kakao_address.AddressLookupError as exc:
            cache[address] = str(exc)
    found = cache[address]
    if isinstance(found, str):
        if ms._is_blank(data.get("latitude")) or ms._is_blank(data.get("longitude")):
            return f"{found} 도로명 주소를 확인해 주세요."
        return None
    for key, value in found.items():
        if ms._is_blank(data.get(key)):
            data[key] = value
    return None


def _missing_required(data: dict) -> list[str]:
    need = set(REQUIRED)
    need |= REQUIRED_MOVING if data.get("media_source") == "MOVING" else REQUIRED_FIXED
    return sorted(k for k in need if ms._is_blank(data.get(k)))


def import_media_xlsx(db: Session, content: bytes) -> dict:
    """엑셀 일괄 등록 — 한글 양식이면 여기서, 예전 영문 양식이면 media_service 로 넘긴다.

    한글 양식은 어드민 폼과 같은 규칙(필수 칸, 음수, 상품 하나 이상)을 지키고, 매체 ID 는 자동 부여한다.
    매체와 상품은 "연결 번호"로 잇는다(다운로드 엑셀은 매체 ID로 이어져 있다).
    매체 ID 칸에 이미 있는 ID가 적힌 줄(다운로드 엑셀을 그대로 올린 경우)은 건너뛴다(수정하지 않음).
    반환: {total, inserted, skipped, failed, errors}. 한 줄이 실패해도 나머지는 계속 넣는다.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="엑셀 파일을 읽을 수 없습니다.")
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="빈 파일입니다.")
    if not is_korean_layout(rows):
        return ms.import_media_xlsx(db, content)

    labels = _header_key_map()
    header = [_clean_label(v) for v in rows[1]]
    columns: list[tuple[int, str, object]] = []
    for idx, label in enumerate(header):
        found = labels.get(label)
        if found and found[0] not in UNUSED_COLUMNS:
            columns.append((idx, *found))
    names = {key: label for _, cols in SECTIONS for key, label, _ in cols}
    # 양식으로 올렸으면 오류 문구도 양식의 칸 이름으로
    names.update({k: label for k, label in _TEMPLATE_LABELS.items() if label in header})
    existing_ids = {r[0] for r in db.query(Media.media_id).all()}
    skip_keys = {"media_id", "created_at", "updated_at"} | _PLAN_DERIVED

    # 매체와 상품을 잇는 칸 — 업로드 양식은 연결 번호, 다운로드 엑셀은 매체 ID.
    link_label = LINK_LABEL if LINK_LABEL in header else "매체 ID"
    link_idx = header.index(link_label) if link_label in header else None
    plans_by_link, plan_errors = _read_plans(wb, link_label)

    inserted = skipped = failed = 0
    errors: list[dict] = []
    seen_links: set[str] = set()
    address_cache: dict[str, dict | str] = {}

    def fail(row_no: int, reason: str) -> None:
        nonlocal failed
        failed += 1
        if len(errors) < ms._IMPORT_MAX_ERRORS:
            errors.append({"row": row_no, "media_id": "", "reason": reason})

    for row_no, row in enumerate(rows[2:], start=3):
        if all(ms._is_blank(v) for v in row):
            continue
        data: dict = {}
        bad = None
        for idx, key, conv in columns:
            raw = row[idx] if idx < len(row) else None
            if key == "media_id":
                if not ms._is_blank(raw) and str(raw).strip() in existing_ids:
                    bad = "skip"
                    break
                continue
            if key in skip_keys:
                continue
            try:
                data[key] = _parse_cell(key, conv, raw)
            except (ValueError, TypeError, json.JSONDecodeError) as exc:
                bad = f"{names[key]}: {exc}"
                break
        if bad == "skip":
            skipped += 1
            # 건너뛴 매체의 상품 줄도 함께 넘긴다("매체가 없다" 오류로 잡지 않게).
            seen_links.add(_link_key(row[link_idx]) if link_idx is not None and link_idx < len(row) else "")
            continue
        if bad:
            fail(row_no, bad)
            continue
        # 매체 칸 오류를 먼저 알리고, 그다음 상품 연결을 본다.
        address_error = _fill_address(data, address_cache)
        if address_error:
            fail(row_no, address_error)
            continue
        missing = _missing_required(data)
        if missing:
            fail(row_no, "필수 칸이 비었습니다: " + ", ".join(names[k] for k in missing))
            continue
        link = _link_key(row[link_idx] if link_idx is not None and link_idx < len(row) else None)
        if not link:
            fail(row_no, f"{link_label}이 비었습니다. 상품 시트와 잇는 번호를 적어 주세요.")
            continue
        if link in seen_links:
            fail(row_no, f"{link_label} {link}이(가) 다른 줄과 겹칩니다.")
            continue
        seen_links.add(link)
        plans = [plan for _, plan in plans_by_link.get(link, [])]
        if not plans:
            fail(row_no, f"상품이 없습니다. 상품 시트에 {link_label} {link}로 상품을 하나 이상 적어 주세요.")
            continue
        ms._apply_spec_unit(data)
        try:
            ms._apply_dooh_rules(data)
            ms._validate_media_values(data)
        except HTTPException as exc:
            fail(row_no, str(exc.detail))
            continue
        data = {k: v for k, v in data.items() if v is not None}
        data["media_id"] = ms._next_media_id(db, data.get("media_source"))
        try:
            with db.begin_nested():
                media = Media(**data)
                db.add(media)
                db.flush()
                ms._sync_plans(db, media, plans)
                ms._apply_plan_aggregates(media)
                db.flush()
            inserted += 1
        except HTTPException as exc:  # 상품 검사(media_service._clean_plan) 실패
            fail(row_no, f"상품 시트({link_label} {link}): {exc.detail}")
        except Exception:
            fail(row_no, "저장 실패")

    # 어느 매체에도 이어지지 않은 상품 줄 — 연결 번호를 잘못 적은 경우.
    for link, items in plans_by_link.items():
        if link not in seen_links:
            for plan_row, _ in items:
                plan_errors.append({
                    "row": plan_row,
                    "media_id": "",
                    "reason": f"상품 시트: {link_label} {link}인 매체가 매체 등록 시트에 없습니다.",
                })
    errors.extend(plan_errors[: max(0, ms._IMPORT_MAX_ERRORS - len(errors))])

    db.commit()
    return {
        "total": inserted + skipped + failed,
        "inserted": inserted,
        "skipped": skipped,
        "failed": failed,
        "errors": errors,
    }
